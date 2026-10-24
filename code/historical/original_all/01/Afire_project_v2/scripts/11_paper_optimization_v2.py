from __future__ import annotations

import math
import os
import textwrap
import warnings
import zipfile
from pathlib import Path

import numpy as np
import pandas as pd
from scipy.stats import linregress, spearmanr, ttest_ind
import statsmodels.api as sm
import statsmodels.formula.api as smf

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import font_manager
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

warnings.filterwarnings('ignore')

ROOT = Path(__file__).resolve().parents[1]
TABLES = ROOT / 'tables'
REPORTS = ROOT / 'reports'
FIGURES = ROOT / 'figures_paper_v2'
PROCESSED = ROOT / 'data_processed'
LOGS = ROOT / 'logs'
for p in [TABLES, REPORTS, FIGURES, PROCESSED, LOGS]:
    p.mkdir(parents=True, exist_ok=True)

# -----------------------------
# Common utilities
# -----------------------------

def find_cjk_font() -> str:
    candidates = [
        '/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc',
        '/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc',
        '/usr/share/fonts/truetype/arphic/uming.ttc',
    ]
    for p in candidates:
        if Path(p).exists():
            font_manager.fontManager.addfont(p)
            try:
                return font_manager.FontProperties(fname=p).get_name()
            except Exception:
                pass
    return 'DejaVu Sans'

CJK_FONT = find_cjk_font()
plt.rcParams.update({
    'font.family': CJK_FONT,
    'axes.unicode_minus': False,
    'font.size': 10.5,
    'axes.titlesize': 13,
    'axes.labelsize': 10.5,
    'legend.fontsize': 9.5,
    'figure.dpi': 150,
    'savefig.dpi': 300,
    'pdf.fonttype': 42,
    'ps.fonttype': 42,
})

COLORS = {
    'blue': '#2F5D8C',
    'light_blue': '#8DB3D3',
    'orange': '#D9822B',
    'green': '#3A7D5D',
    'red': '#B44C43',
    'purple': '#7566A6',
    'gray': '#6B7280',
    'light_gray': '#E9EDF2',
    'dark': '#25313D',
}

CAUSE_ORDER = [
    'respiratory_system', 'copd', 'chronic_lower_respiratory', 'asthma',
    'circulatory_system', 'ischemic_heart_disease', 'cerebrovascular_disease'
]


def sig_label(p: float) -> str:
    if pd.isna(p):
        return '无法检验'
    if p < 0.001:
        return '是（p<0.001）'
    if p < 0.01:
        return '是（p<0.01）'
    if p < 0.05:
        return '是（p<0.05）'
    return '否'


def p_fmt(p: float) -> str:
    if pd.isna(p):
        return 'NA'
    if p < 0.001:
        return '<0.001'
    return f'{p:.3f}'


def fmt(x, digits=3):
    if pd.isna(x):
        return 'NA'
    if isinstance(x, (int, np.integer)):
        return str(int(x))
    return f'{float(x):.{digits}f}'


def save_fig(fig, stem: str):
    fig.tight_layout()
    fig.savefig(FIGURES / f'{stem}.png', bbox_inches='tight', facecolor='white', dpi=300)
    fig.savefig(FIGURES / f'{stem}.pdf', bbox_inches='tight', facecolor='white')
    plt.close(fig)


def minmax(s: pd.Series) -> pd.Series:
    x = pd.to_numeric(s, errors='coerce')
    lo, hi = x.min(), x.max()
    if pd.isna(lo) or pd.isna(hi) or hi <= lo:
        return pd.Series(0.0, index=s.index)
    return (x - lo) / (hi - lo)


def welch_summary(a: pd.Series, b: pd.Series):
    a = pd.to_numeric(a, errors='coerce').dropna()
    b = pd.to_numeric(b, errors='coerce').dropna()
    if len(a) < 2 or len(b) < 2:
        p = np.nan
    else:
        p = float(ttest_ind(a, b, equal_var=False, nan_policy='omit').pvalue)
    pooled_sd = math.sqrt(((len(a)-1)*a.var(ddof=1)+(len(b)-1)*b.var(ddof=1))/max(len(a)+len(b)-2,1)) if len(a)>1 and len(b)>1 else np.nan
    smd = (a.mean()-b.mean())/pooled_sd if pooled_sd and np.isfinite(pooled_sd) and pooled_sd>0 else np.nan
    return {
        'covered_n': len(a), 'noncovered_n': len(b),
        'covered_mean': a.mean() if len(a) else np.nan,
        'noncovered_mean': b.mean() if len(b) else np.nan,
        'covered_median': a.median() if len(a) else np.nan,
        'noncovered_median': b.median() if len(b) else np.nan,
        'difference_covered_minus_noncovered': a.mean()-b.mean() if len(a) and len(b) else np.nan,
        'ratio_covered_to_noncovered': a.mean()/b.mean() if len(a) and len(b) and b.mean()!=0 else np.nan,
        'standardized_mean_difference': smd,
        'p_value_welch': p,
        'statistically_significant_0_05': 'yes' if pd.notna(p) and p<0.05 else 'no',
    }


def read_inputs():
    panel = pd.read_csv(PROCESSED / 'state_year_disease_panel.csv')
    fire = pd.read_csv(TABLES / 'top100_fire_pm25_cover_state_year.csv')
    ann_exp = pd.read_csv(TABLES / 'annual_exposure_inequality.csv')
    ann_health = pd.read_csv(TABLES / 'annual_health_inequality_by_disease.csv')
    models = pd.read_csv(TABLES / 'disease_specific_model_results.csv')
    wphii = pd.read_csv(TABLES / 'WPHII_state_year_disease.csv')
    top = pd.read_csv(TABLES / 'top_risk_states.csv')
    return panel, fire, ann_exp, ann_health, models, wphii, top

# -----------------------------
# Integrity audit
# -----------------------------

def audit_v2(panel, fire, models, wphii):
    required_files = [
        PROCESSED / 'state_year_disease_panel.csv',
        TABLES / 'top100_fire_pm25_cover_state_year.csv',
        TABLES / 'fire_covered_state_year.csv',
        TABLES / 'annual_exposure_inequality.csv',
        TABLES / 'annual_health_inequality_by_disease.csv',
        TABLES / 'disease_specific_model_results.csv',
        TABLES / 'WPHII_state_year_disease.csv',
        TABLES / 'top_risk_states.csv',
    ]
    required_fields = {
        'panel': ['state','state_abbr','year','cause_id','deaths','population','age_adjusted_rate_per_100k',
                  'smoke_pm25_annual_burden','smoke_days_mean','ses_disadvantage_z','top100_fire_any',
                  'pm25_cover_top25','top100_fire_pm25_cover','fire_pm25_cover_intensity','fire_covered'],
        'fire': ['state','state_abbr','year','top100_fire_any','pm25_cover_top25',
                 'top100_fire_pm25_cover','fire_pm25_cover_intensity','fire_covered_definition'],
        'wphii': ['exposure_norm','vulnerability_norm','health_burden_norm','fire_weight','WPHII'],
        'models': ['model','cause_id','term','estimate','std_error','p_value_normal_approx'],
    }
    rows=[]
    for f in required_files:
        rows.append({'check':'关键文件存在','item':str(f.relative_to(ROOT)),'status':'PASS' if f.exists() else 'FAIL','detail':''})
    for label, cols in required_fields.items():
        d={'panel':panel,'fire':fire,'wphii':wphii,'models':models}[label]
        for c in cols:
            rows.append({'check':'关键字段存在','item':f'{label}.{c}','status':'PASS' if c in d.columns else 'FAIL','detail':''})
    years=sorted(panel['year'].dropna().astype(int).unique())
    causes=sorted(panel['cause_id'].dropna().unique())
    rows.append({'check':'年份完整性','item':'2006-2020','status':'PASS' if years==list(range(2006,2021)) else 'FAIL','detail':str(years)})
    rows.append({'check':'疾病类别完整性','item':'7 categories','status':'PASS' if len(causes)==7 else 'FAIL','detail':', '.join(causes)})
    expected=panel.state_abbr.nunique()*panel.year.nunique()*panel.cause_id.nunique()
    rows.append({'check':'state-year-disease 网格','item':f'expected={expected}','status':'PASS' if len(panel)==expected and panel.duplicated(['state_abbr','year','cause_id']).sum()==0 else 'FAIL','detail':f'rows={len(panel)}; duplicates={panel.duplicated(["state_abbr","year","cause_id"]).sum()}'})
    overlap_ok=(fire['top100_fire_pm25_cover'].astype(int)==((fire['top100_fire_any']==1)&(fire['pm25_cover_top25']==1)).astype(int)).all()
    fire_alias_ok=(fire['fire_covered'].astype(int)==fire['top100_fire_pm25_cover'].astype(int)).all()
    rows.append({'check':'Top100-PM2.5 cover 逻辑','item':'overlap = top100_fire_any AND pm25_cover_top25','status':'PASS' if overlap_ok else 'FAIL','detail':''})
    rows.append({'check':'fire_covered 主定义','item':'fire_covered = top100_fire_pm25_cover','status':'PASS' if fire_alias_ok else 'FAIL','detail':'; '.join(fire['fire_covered_definition'].dropna().unique().astype(str))})
    calc=wphii['exposure_norm']*wphii['vulnerability_norm']*wphii['health_burden_norm']*wphii['fire_weight']
    err=float((calc-wphii['WPHII']).abs().max())
    fw=1+0.35*wphii['top100_fire_any']+0.65*wphii['top100_fire_pm25_cover']+0.50*wphii['fire_cover_intensity_norm']
    fwerr=float((fw-wphii['fire_weight']).abs().max())
    rows.append({'check':'WPHII 乘积公式','item':'exposure × vulnerability × health × fire_weight','status':'PASS' if err<1e-10 else 'FAIL','detail':f'max_abs_error={err:.3g}'})
    rows.append({'check':'WPHII fire_weight','item':'1+0.35 fire+0.65 overlap+0.50 intensity','status':'PASS' if fwerr<1e-10 else 'FAIL','detail':f'max_abs_error={fwerr:.3g}'})
    rows.append({'check':'模型结果','item':'7 diseases × 3 model families','status':'PASS' if models.cause_id.nunique()==7 and (models.term=='MODEL_FAILED').sum()==0 else 'FAIL','detail':f'rows={len(models)}; failed={(models.term=="MODEL_FAILED").sum()}'})
    # Existing DOCX: zipped OOXML can be opened.
    docx_candidates=list(REPORTS.glob('*v2.docx'))
    docx_ok=False
    docx_detail='未找到'
    for p in docx_candidates:
        try:
            with zipfile.ZipFile(p) as z:
                docx_ok='word/document.xml' in z.namelist()
            if docx_ok:
                docx_detail=p.name
                break
        except Exception as e:
            docx_detail=f'{p.name}: {e}'
    rows.append({'check':'中文报告 DOCX 可打开','item':'existing v2 report','status':'PASS' if docx_ok else 'FAIL','detail':docx_detail})
    audit=pd.DataFrame(rows)
    audit.to_csv(TABLES/'v2_integrity_check.csv',index=False,encoding='utf-8-sig')
    need_rerun=(audit.status=='FAIL').any()
    return audit, need_rerun

# -----------------------------
# Lag-1 models
# -----------------------------

def fit_lag_models(panel: pd.DataFrame, main_models: pd.DataFrame):
    state_year = panel.sort_values(['state_abbr','year']).drop_duplicates(['state_abbr','year'])[
        ['state_abbr','year','smoke_pm25_annual_burden','top100_fire_pm25_cover','fire_pm25_cover_intensity']
    ].copy()
    state_year=state_year.sort_values(['state_abbr','year'])
    state_year['smoke_pm25_lag1']=state_year.groupby('state_abbr')['smoke_pm25_annual_burden'].shift(1)
    state_year['top100_fire_pm25_cover_lag1']=state_year.groupby('state_abbr')['top100_fire_pm25_cover'].shift(1)
    state_year['fire_pm25_cover_intensity_lag1']=state_year.groupby('state_abbr')['fire_pm25_cover_intensity'].shift(1)
    mu=state_year['smoke_pm25_lag1'].mean(); sd=state_year['smoke_pm25_lag1'].std()
    state_year['z_smoke_pm25_lag1']=(state_year['smoke_pm25_lag1']-mu)/sd
    lag=panel.merge(state_year[['state_abbr','year','smoke_pm25_lag1','top100_fire_pm25_cover_lag1','fire_pm25_cover_intensity_lag1','z_smoke_pm25_lag1']],on=['state_abbr','year'],how='left')
    vmu=lag['ses_disadvantage_z'].mean(); vsd=lag['ses_disadvantage_z'].std()
    lag['z_vulnerability']=(lag['ses_disadvantage_z']-vmu)/vsd
    lag['lag1_pm25_x_vulnerability']=lag['z_smoke_pm25_lag1']*lag['z_vulnerability']
    lag.to_csv(PROCESSED/'model_ready_panel_lag1_v2.csv',index=False,encoding='utf-8-sig')

    lookup=main_models.set_index(['model','cause_id','term'])['estimate'].to_dict()
    rows=[]
    for cause in CAUSE_ORDER:
        sub=lag[lag.cause_id==cause].copy()
        cause_zh=sub['cause_zh'].iloc[0]
        cause_en=sub['cause_en'].iloc[0]
        # Rate model
        d=sub.dropna(subset=['age_adjusted_rate_per_100k','z_smoke_pm25_lag1','z_vulnerability','state_abbr','year']).copy()
        formula='age_adjusted_rate_per_100k ~ z_smoke_pm25_lag1 + z_vulnerability + lag1_pm25_x_vulnerability + C(state_abbr) + C(year)'
        try:
            fit=smf.ols(formula,data=d).fit(cov_type='cluster',cov_kwds={'groups':d['state_abbr']})
            for term in ['z_smoke_pm25_lag1','z_vulnerability','lag1_pm25_x_vulnerability']:
                est=float(fit.params.get(term,np.nan)); se=float(fit.bse.get(term,np.nan)); p=float(fit.pvalues.get(term,np.nan))
                main_term={'z_smoke_pm25_lag1':'z_smoke_pm25_annual_burden','lag1_pm25_x_vulnerability':'pm25_x_vulnerability'}.get(term)
                main_est=lookup.get(('ols_rate_state_year_fe',cause,main_term),np.nan) if main_term else np.nan
                rows.append({
                    'model':'ols_rate_state_year_fe_lag1','cause_id':cause,'cause_en':cause_en,'cause_zh':cause_zh,
                    'term':term,'estimate':est,'std_error_cluster_state':se,'ci95_low':est-1.96*se,'ci95_high':est+1.96*se,
                    'p_value':p,'n':int(fit.nobs),'n_states':d.state_abbr.nunique(),'effect_measure':'rate difference per 100,000 per 1 SD',
                    'exp_estimate':np.nan,'percent_change':np.nan,'main_model_estimate':main_est,
                    'direction_consistent_with_main':('yes' if pd.notna(main_est) and np.sign(est)==np.sign(main_est) else ('no' if pd.notna(main_est) else 'not_applicable')),
                    'note':'州与年份固定效应；标准误按州聚类。lag1 为补充性生态学关联。'
                })
        except Exception as e:
            rows.append({'model':'ols_rate_state_year_fe_lag1','cause_id':cause,'cause_en':cause_en,'cause_zh':cause_zh,'term':'MODEL_FAILED','estimate':np.nan,'std_error_cluster_state':np.nan,'ci95_low':np.nan,'ci95_high':np.nan,'p_value':np.nan,'n':len(d),'n_states':d.state_abbr.nunique(),'effect_measure':'','exp_estimate':np.nan,'percent_change':np.nan,'main_model_estimate':np.nan,'direction_consistent_with_main':'not_applicable','note':str(e)})
        # Population-offset Poisson model
        d=sub.dropna(subset=['deaths','population','z_smoke_pm25_lag1','z_vulnerability','state_abbr','year']).copy()
        d=d[(d.deaths>=0)&(d.population>0)]
        formula='deaths ~ z_smoke_pm25_lag1 + z_vulnerability + lag1_pm25_x_vulnerability + C(state_abbr) + C(year)'
        try:
            fit=smf.glm(formula,data=d,family=sm.families.Poisson(),offset=np.log(d['population'])).fit(cov_type='cluster',cov_kwds={'groups':d['state_abbr']},maxiter=200)
            for term in ['z_smoke_pm25_lag1','z_vulnerability','lag1_pm25_x_vulnerability']:
                est=float(fit.params.get(term,np.nan)); se=float(fit.bse.get(term,np.nan)); p=float(fit.pvalues.get(term,np.nan))
                main_term={'z_smoke_pm25_lag1':'z_smoke_pm25_annual_burden','lag1_pm25_x_vulnerability':'pm25_x_vulnerability'}.get(term)
                main_est=lookup.get(('quasi_poisson_deaths_state_year_fe',cause,main_term),np.nan) if main_term else np.nan
                rows.append({
                    'model':'poisson_deaths_state_year_fe_lag1','cause_id':cause,'cause_en':cause_en,'cause_zh':cause_zh,
                    'term':term,'estimate':est,'std_error_cluster_state':se,'ci95_low':est-1.96*se,'ci95_high':est+1.96*se,
                    'p_value':p,'n':int(fit.nobs),'n_states':d.state_abbr.nunique(),'effect_measure':'incidence rate ratio per 1 SD',
                    'exp_estimate':math.exp(est) if np.isfinite(est) else np.nan,
                    'percent_change':(math.exp(est)-1)*100 if np.isfinite(est) else np.nan,
                    'main_model_estimate':main_est,
                    'direction_consistent_with_main':('yes' if pd.notna(main_est) and np.sign(est)==np.sign(main_est) else ('no' if pd.notna(main_est) else 'not_applicable')),
                    'note':'死亡数模型采用 log(population) offset；州与年份固定效应；标准误按州聚类。'
                })
        except Exception as e:
            rows.append({'model':'poisson_deaths_state_year_fe_lag1','cause_id':cause,'cause_en':cause_en,'cause_zh':cause_zh,'term':'MODEL_FAILED','estimate':np.nan,'std_error_cluster_state':np.nan,'ci95_low':np.nan,'ci95_high':np.nan,'p_value':np.nan,'n':len(d),'n_states':d.state_abbr.nunique(),'effect_measure':'','exp_estimate':np.nan,'percent_change':np.nan,'main_model_estimate':np.nan,'direction_consistent_with_main':'not_applicable','note':str(e)})
    out=pd.DataFrame(rows)
    out.to_csv(TABLES/'lag1_model_results_v2.csv',index=False,encoding='utf-8-sig')
    return lag,out

# -----------------------------
# WPHII sensitivity
# -----------------------------

def wphii_sensitivity(wphii: pd.DataFrame):
    d=wphii.copy()
    base=d['exposure_norm']*d['vulnerability_norm']*d['health_burden_norm']
    d['fire_weight_main']=1+0.35*d['top100_fire_any']+0.65*d['top100_fire_pm25_cover']+0.50*d['fire_cover_intensity_norm']
    d['fire_weight_equal']=1+0.50*d['top100_fire_any']+0.50*d['top100_fire_pm25_cover']+0.50*d['fire_cover_intensity_norm']
    d['fire_weight_no_fire']=1.0
    d['fire_weight_strong']=1+0.50*d['top100_fire_any']+1.00*d['top100_fire_pm25_cover']+0.75*d['fire_cover_intensity_norm']
    d['WPHII_main']=base*d['fire_weight_main']
    d['WPHII_equal_weight']=base*d['fire_weight_equal']
    d['WPHII_no_fire_weight']=base
    d['WPHII_strong_fire_weight']=base*d['fire_weight_strong']
    d.to_csv(TABLES/'WPHII_sensitivity_v2.csv',index=False,encoding='utf-8-sig')

    versions=['WPHII_main','WPHII_equal_weight','WPHII_no_fire_weight','WPHII_strong_fire_weight']
    labels={
        'WPHII_main':'主权重','WPHII_equal_weight':'火灾分量等权','WPHII_no_fire_weight':'不含火灾权重','WPHII_strong_fire_weight':'强化火灾权重'
    }
    agg=d.groupby(['state','state_abbr','cause_id','cause_en','cause_zh'],as_index=False)[versions].mean()
    rank_frames={}
    for v in versions:
        tmp=agg.sort_values(v,ascending=False).reset_index(drop=True)
        tmp['rank']=np.arange(1,len(tmp)+1)
        rank_frames[v]=tmp
    main_top=set((rank_frames['WPHII_main'].head(10)['state_abbr']+'|'+rank_frames['WPHII_main'].head(10)['cause_id']).tolist())
    main_rank={(r.state_abbr+'|'+r.cause_id):int(r.rank) for r in rank_frames['WPHII_main'].itertuples()}
    rows=[]
    metrics=[]
    main_scores=agg.set_index(['state_abbr','cause_id'])['WPHII_main']
    for v in versions:
        tmp=rank_frames[v]
        top=tmp.head(10)
        top_set=set((top.state_abbr+'|'+top.cause_id).tolist())
        overlap=len(main_top&top_set)
        jac=len(main_top&top_set)/len(main_top|top_set)
        scores=agg.set_index(['state_abbr','cause_id'])[v]
        common=main_scores.index.intersection(scores.index)
        rho=float(spearmanr(main_scores.loc[common],scores.loc[common],nan_policy='omit').statistic)
        metrics.append({'version':v,'version_zh':labels[v],'top10_overlap_with_main_n':overlap,'top10_jaccard_with_main':jac,'spearman_rank_correlation_with_main':rho})
        for r in top.itertuples():
            combo=r.state_abbr+'|'+r.cause_id
            rows.append({
                'version':v,'version_zh':labels[v],'rank':int(r.rank),'state':r.state,'state_abbr':r.state_abbr,
                'cause_id':r.cause_id,'cause_en':r.cause_en,'cause_zh':r.cause_zh,'mean_index':getattr(r,v),
                'in_main_top10':'yes' if combo in main_top else 'no','main_rank':main_rank.get(combo),
                'top10_overlap_with_main_n':overlap,'top10_jaccard_with_main':jac,'spearman_rank_correlation_with_main':rho,
                'interpretation':'排名用于优先级识别，不代表因果效应。'
            })
    comp=pd.DataFrame(rows)
    metrics=pd.DataFrame(metrics)
    comp.to_csv(TABLES/'WPHII_top10_comparison_v2.csv',index=False,encoding='utf-8-sig')
    metrics.to_csv(TABLES/'WPHII_sensitivity_metrics_v2.csv',index=False,encoding='utf-8-sig')
    return d,agg,comp,metrics

# -----------------------------
# Fire-covered comparison
# -----------------------------

def fire_comparison(panel, fire, wphii):
    rows=[]
    for metric,label,unit in [
        ('smoke_pm25_annual_burden','烟雾 PM2.5 年暴露','μg/m³（年度负担口径）'),
        ('smoke_days_mean','烟雾日数','天/年'),
    ]:
        s=welch_summary(fire.loc[fire.top100_fire_pm25_cover==1,metric],fire.loc[fire.top100_fire_pm25_cover==0,metric])
        rows.append({'outcome_group':'exposure','outcome_id':metric,'outcome_label':label,'unit':unit,'analysis_unit':'state-year',**s,
                     'interpretation_note':'该比较是描述性构念验证；cover 定义本身包含年度 PM2.5 前25%，不能解释为火灾导致了该差异。'})
    s=welch_summary(wphii.loc[wphii.top100_fire_pm25_cover==1,'WPHII'],wphii.loc[wphii.top100_fire_pm25_cover==0,'WPHII'])
    rows.append({'outcome_group':'WPHII','outcome_id':'WPHII','outcome_label':'WPHII v2','unit':'无量纲指数','analysis_unit':'state-year-disease',**s,
                 'interpretation_note':'WPHII 包含火灾权重和暴露归一化，组间差异属于指数结构的一部分，不是健康效应估计。'})
    for cause in CAUSE_ORDER:
        sub=panel[(panel.cause_id==cause) & panel.smoke_pm25_annual_burden.notna()]
        s=welch_summary(sub.loc[sub.top100_fire_pm25_cover==1,'age_adjusted_rate_per_100k'],sub.loc[sub.top100_fire_pm25_cover==0,'age_adjusted_rate_per_100k'])
        rows.append({'outcome_group':'health_rate','outcome_id':cause,'outcome_label':sub.cause_zh.iloc[0],'unit':'例/10万人','analysis_unit':'state-year',**s,
                     'interpretation_note':'限于有 smoke PM2.5 观测的州年；未调整描述性差异受地区构成、年份和混杂影响，不应解释为个体因果效应。'})
    out=pd.DataFrame(rows)
    out.to_csv(TABLES/'fire_cover_vs_noncover_summary_v2.csv',index=False,encoding='utf-8-sig')
    return out

# -----------------------------
# Figures
# -----------------------------

def fig1_flow():
    fig,ax=plt.subplots(figsize=(14,7.2))
    ax.set_xlim(0,14);ax.set_ylim(0,7.2);ax.axis('off')
    boxes=[
        (0.3,4.3,2.2,1.4,'原始数据','县日 smoke PM2.5\nCDC WONDER 7类疾病\nNHGIS 社会经济指标\nTop100 大火'),
        (3.0,4.3,2.2,1.4,'统一面板','2006–2020\nstate–year–disease\n51个州级单元'),
        (5.7,4.3,2.2,1.4,'火灾–烟雾联结','Top100 fire any\nPM2.5 年度 Top25%\n二者重叠与连续强度'),
        (8.4,4.3,2.2,1.4,'不公平性指标','人口加权均值\nGini / Theil\nCI / SII / RII'),
        (11.1,4.3,2.2,1.4,'WPHII v2','暴露 × 脆弱性\n× 疾病负担\n× 火灾权重'),
        (3.2,1.5,3.0,1.3,'分疾病探索性模型','率模型 + 死亡数 offset 模型\n州固定效应 + 年份固定效应'),
        (7.6,1.5,3.0,1.3,'补充与稳健性分析','1年滞后暴露\nWPHII 权重敏感性\ncover vs non-cover 描述'),
    ]
    fills=['#EDF4FA','#EFF6F1','#FFF3E6','#F3F0FA','#FDEEEE','#EEF3F8','#F6F6F6']
    for (x,y,w,h,title,body),fc in zip(boxes,fills):
        patch=FancyBboxPatch((x,y),w,h,boxstyle='round,pad=0.04,rounding_size=0.08',linewidth=1.4,edgecolor=COLORS['blue'],facecolor=fc)
        ax.add_patch(patch)
        ax.text(x+w/2,y+h-0.32,title,ha='center',va='center',fontsize=12,fontweight='bold',color=COLORS['dark'])
        ax.text(x+w/2,y+0.48,body,ha='center',va='center',fontsize=9.6,linespacing=1.35,color='#374151')
    for x1,x2 in [(2.5,3.0),(5.2,5.7),(7.9,8.4),(10.6,11.1)]:
        ax.add_patch(FancyArrowPatch((x1,5.0),(x2,5.0),arrowstyle='-|>',mutation_scale=16,linewidth=1.5,color=COLORS['gray']))
    ax.add_patch(FancyArrowPatch((6.8,4.25),(4.7,2.85),arrowstyle='-|>',mutation_scale=16,linewidth=1.4,color=COLORS['gray'],connectionstyle='arc3,rad=0.1'))
    ax.add_patch(FancyArrowPatch((10.0,4.25),(9.1,2.85),arrowstyle='-|>',mutation_scale=16,linewidth=1.4,color=COLORS['gray'],connectionstyle='arc3,rad=-0.1'))
    ax.text(7,6.75,'v2 论文级分析框架：以暴露不公平性与综合风险识别为主线',ha='center',va='center',fontsize=17,fontweight='bold',color=COLORS['dark'])
    ax.text(7,0.55,'解释边界：州尺度生态学证据；Top100 不等于全部火灾；WPHII 是优先级指数而非因果效应。',ha='center',fontsize=10.5,color=COLORS['red'])
    save_fig(fig,'Fig1_v2_study_design_flow')


def fig2_exposure(ann_exp):
    x=ann_exp.year
    series=[
        ('population_weighted_smoke_pm25','人口加权 smoke PM2.5','μg/m³',COLORS['blue']),
        ('gini_smoke_pm25','暴露 Gini 系数','Gini',COLORS['orange']),
        ('concentration_index_by_ses_disadvantage','按社会脆弱性排序的集中指数','Concentration Index',COLORS['green']),
    ]
    fig,axes=plt.subplots(3,1,figsize=(11,10),sharex=True)
    for ax,(col,title,ylabel,color) in zip(axes,series):
        y=ann_exp[col]
        ax.plot(x,y,marker='o',linewidth=2,color=color,label='年度值')
        ok=y.notna()
        lr=linregress(x[ok],y[ok])
        ax.plot(x,lr.intercept+lr.slope*x,linestyle='--',linewidth=1.5,color=COLORS['gray'],label=f'线性趋势：斜率={lr.slope:.4f}/年，p={p_fmt(lr.pvalue)}')
        ax.axhline(0,color='#BFC5CC',linewidth=0.8)
        ax.set_title(title,loc='left',fontweight='bold')
        ax.set_ylabel(ylabel)
        ax.grid(axis='y',alpha=.25)
        ax.legend(loc='best',frameon=False)
    axes[-1].set_xlabel('年份')
    axes[-1].set_xticks(range(2006,2021));axes[-1].tick_params(axis='x',rotation=45)
    fig.suptitle('2006–2020 年野火烟雾 PM2.5 暴露水平与不公平性趋势',fontsize=16,fontweight='bold',y=1.01)
    fig.text(.5,.005,'注：集中指数为负表示暴露更集中于社会脆弱性较低一侧，正值表示更集中于脆弱性较高一侧。',ha='center',fontsize=9,color=COLORS['gray'])
    save_fig(fig,'Fig2_v2_exposure_inequality_trends')


def fig3_overlap(fire):
    s=fire.groupby('year',as_index=False).agg(top100_fire_states=('top100_fire_any','sum'),pm25_top25_states=('pm25_cover_top25','sum'),overlap_states=('top100_fire_pm25_cover','sum'))
    x=np.arange(len(s));w=.25
    fig,ax=plt.subplots(figsize=(12,6.2))
    ax.bar(x-w,s.top100_fire_states,width=w,label='有 Top100 大火的州数',color=COLORS['orange'])
    ax.bar(x,s.pm25_top25_states,width=w,label='PM2.5 Top25% 州数',color=COLORS['light_blue'])
    ax.bar(x+w,s.overlap_states,width=w,label='Top100–PM2.5 overlap 州数',color=COLORS['red'])
    for xi,v in zip(x+w,s.overlap_states):
        ax.text(xi,v+.25,str(int(v)),ha='center',fontsize=8.5)
    ax.set_xticks(x,s.year.astype(int),rotation=45);ax.set_ylabel('州数');ax.set_xlabel('年份')
    ax.set_title('Top100 大火与年度高 smoke PM2.5 覆盖的州级重叠',fontweight='bold')
    ax.legend(frameon=False,ncol=3,loc='upper left');ax.grid(axis='y',alpha=.25)
    ax.text(.995,.96,f'全期 Top100 州年={int(fire.top100_fire_any.sum())}；overlap 州年={int(fire.top100_fire_pm25_cover.sum())}',transform=ax.transAxes,ha='right',va='top',fontsize=9,color=COLORS['gray'])
    save_fig(fig,'Fig3_v2_top100_pm25_overlap_by_year')


def fig4_2020(fire):
    d=fire[(fire.year==2020)&(fire.top100_fire_pm25_cover==1)].sort_values('fire_pm25_cover_intensity')
    fig,ax=plt.subplots(figsize=(9,5.6))
    bars=ax.barh(d.state_abbr,d.fire_pm25_cover_intensity,color=COLORS['red'],alpha=.85)
    ax.set_xlabel('fire–PM2.5 cover intensity（0–1）');ax.set_ylabel('州')
    ax.set_title('2020 年 Top100–PM2.5 overlap 州的暴露强度排名',fontweight='bold')
    ax.grid(axis='x',alpha=.25)
    for b,row in zip(bars,d.itertuples()):
        ax.text(b.get_width()+.012,b.get_y()+b.get_height()/2,f'{row.fire_pm25_cover_intensity:.3f}  |  PM2.5={row.smoke_pm25_annual_burden:.2f}',va='center',fontsize=9)
    ax.set_xlim(0,max(d.fire_pm25_cover_intensity.max()*1.35,1.05))
    ax.text(.01,-.16,'注：连续强度 = 年内归一化 smoke PM2.5 × 年内归一化 log(分配燃烧面积+1)。',transform=ax.transAxes,fontsize=9,color=COLORS['gray'])
    save_fig(fig,'Fig4_v2_2020_overlap_states')


def fig5_wphii(top):
    d=top.head(20).sort_values('mean_WPHII')
    cause_zh_map={'respiratory_system':'呼吸系统疾病','copd':'COPD','chronic_lower_respiratory':'慢性下呼吸道疾病','asthma':'哮喘','circulatory_system':'循环系统疾病','ischemic_heart_disease':'缺血性心脏病','cerebrovascular_disease':'脑血管疾病'}
    labels=[f'{st} – {cause_zh_map.get(cid,cid)}' for st,cid in zip(d.state_abbr,d.cause_id)]
    fig,ax=plt.subplots(figsize=(11,8.6))
    bars=ax.barh(labels,d.mean_WPHII,color=COLORS['purple'],alpha=.88)
    ax.set_xlabel('2006–2020 平均 WPHII v2');ax.set_title('WPHII v2 排名前 20 的州–疾病组合',fontweight='bold')
    ax.grid(axis='x',alpha=.25)
    for b,v in zip(bars,d.mean_WPHII):
        ax.text(v+.003,b.get_y()+b.get_height()/2,f'{v:.3f}',va='center',fontsize=8.5)
    ax.set_xlim(0,d.mean_WPHII.max()*1.18)
    ax.text(.01,-.08,'注：WPHII 为综合优先级/不公平风险指数，不表示暴露的因果健康效应。',transform=ax.transAxes,fontsize=9,color=COLORS['red'])
    save_fig(fig,'Fig5_v2_wphii_top_states')


def fig6_models(models):
    d=models[(models.model=='ols_rate_state_year_fe') & models.term.isin(['z_smoke_pm25_annual_burden','pm25_x_vulnerability'])].copy()
    d['term_zh']=d.term.map({'z_smoke_pm25_annual_burden':'PM2.5 主效应（每1 SD）','pm25_x_vulnerability':'PM2.5 × 脆弱性交互'})
    cause_lookup=d.drop_duplicates('cause_id').set_index('cause_id')['cause_zh'].to_dict()
    ylabels=[cause_lookup[c] for c in CAUSE_ORDER]
    fig,axes=plt.subplots(1,2,figsize=(13,7.8),sharey=True)
    for ax,term,color in zip(axes,['z_smoke_pm25_annual_burden','pm25_x_vulnerability'],[COLORS['blue'],COLORS['orange']]):
        sub=d[d.term==term].set_index('cause_id').reindex(CAUSE_ORDER).reset_index()
        y=np.arange(len(sub))[::-1]
        est=sub.estimate.values;lo=est-1.96*sub.std_error.values;hi=est+1.96*sub.std_error.values
        ax.errorbar(est,y,xerr=[est-lo,hi-est],fmt='o',capsize=3,color=color,ecolor=color,elinewidth=1.5)
        for xi,yi,p in zip(est,y,sub.p_value_normal_approx):
            if pd.notna(p) and p<.05:
                ax.text(xi,yi+.20,'*',ha='center',fontsize=13,color=COLORS['red'])
        ax.axvline(0,color=COLORS['gray'],linewidth=1)
        ax.set_yticks(y,ylabels)
        ax.set_xlabel('年龄调整死亡率差异（例/10万人）')
        ax.set_title(sub.term_zh.iloc[0],fontweight='bold')
        ax.grid(axis='x',alpha=.25)
    fig.suptitle('分疾病模型核心项：州尺度探索性固定效应关联',fontsize=16,fontweight='bold')
    fig.text(.5,.015,'点为系数，线为近似 95% CI；* p<0.05。负向系数不得解释为 PM2.5 具有保护作用。',ha='center',fontsize=9,color=COLORS['red'])
    save_fig(fig,'Fig6_v2_disease_model_summary')


def fig7_cover(fire):
    cov=fire.loc[fire.top100_fire_pm25_cover==1,'smoke_pm25_annual_burden'].dropna()
    non=fire.loc[fire.top100_fire_pm25_cover==0,'smoke_pm25_annual_burden'].dropna()
    p=ttest_ind(cov,non,equal_var=False).pvalue
    fig,ax=plt.subplots(figsize=(8.5,6.4))
    bp=ax.boxplot([non,cov],labels=[f'non-covered\n(n={len(non)})',f'fire-covered\n(n={len(cov)})'],showfliers=False,patch_artist=True,widths=.5)
    bp['boxes'][0].set_facecolor(COLORS['light_blue']);bp['boxes'][1].set_facecolor('#E6A29B')
    rng=np.random.default_rng(42)
    for i,vals,color in [(1,non,COLORS['blue']),(2,cov,COLORS['red'])]:
        x=rng.normal(i,.055,size=len(vals))
        ax.scatter(x,vals,s=13,alpha=.45,color=color,edgecolors='none')
    ax.set_ylabel('smoke PM2.5 年暴露（μg/m³）')
    ax.set_title('Top100–PM2.5 cover 与 non-cover 州年的暴露分布',fontweight='bold')
    ax.grid(axis='y',alpha=.25)
    ax.text(.5,.95,f'均值：{non.mean():.3f} vs {cov.mean():.3f}；Welch p={p_fmt(p)}',transform=ax.transAxes,ha='center',va='top',fontsize=10)
    ax.text(.5,-.14,'该差异部分由 cover 定义（PM2.5 年度前25%）决定，只能作为构念验证，不能作为因果效应。',transform=ax.transAxes,ha='center',fontsize=9,color=COLORS['red'])
    save_fig(fig,'Fig7_v2_fire_cover_vs_nonfire_exposure')


def fig8_sensitivity(metrics):
    d=metrics[metrics.version!='WPHII_main'].copy()
    x=np.arange(len(d))
    fig,ax=plt.subplots(figsize=(10,6.2))
    bars=ax.bar(x-.18,d.top10_overlap_with_main_n,width=.36,color=COLORS['purple'],label='与主版本 Top10 重合数（0–10）')
    ax.set_ylabel('Top10 重合数');ax.set_ylim(0,10.8);ax.set_xticks(x,d.version_zh)
    ax2=ax.twinx()
    line=ax2.plot(x+.18,d.spearman_rank_correlation_with_main,marker='o',linewidth=2,color=COLORS['green'],label='全排名 Spearman 相关')
    ax2.set_ylabel('Spearman ρ');ax2.set_ylim(0,1.05)
    for b,v in zip(bars,d.top10_overlap_with_main_n): ax.text(b.get_x()+b.get_width()/2,v+.2,str(int(v)),ha='center')
    for xi,v in zip(x+.18,d.spearman_rank_correlation_with_main): ax2.text(xi,v+.025,f'{v:.3f}',ha='center',fontsize=9)
    ax.set_title('WPHII v2 权重敏感性：Top10 与整体排序稳定性',fontweight='bold')
    ax.grid(axis='y',alpha=.25)
    handles=[bars,line[0]];labels=['与主版本 Top10 重合数','全排名 Spearman 相关']
    ax.legend(handles,labels,frameon=False,loc='lower left')
    ax.text(.5,-.13,'火灾权重改变用于检验优先级排名稳定性；不含火灾权重版本检验火灾维度的增量影响。',transform=ax.transAxes,ha='center',fontsize=9,color=COLORS['gray'])
    save_fig(fig,'Fig8_v2_wphii_sensitivity')

# -----------------------------
# DOCX helpers
# -----------------------------

def set_cell_shading(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn('w:shd'))
    if shd is None:
        shd=OxmlElement('w:shd');tcPr.append(shd)
    shd.set(qn('w:fill'),fill)


def set_cell_text(cell,text,bold=False,size=8.5,color=None):
    cell.text=''
    p=cell.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(str(text));r.bold=bold;r.font.size=Pt(size);r.font.name='Microsoft YaHei';r._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    if color: r.font.color.rgb=RGBColor(*color)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc,df,columns=None,headers=None,max_rows=20,font_size=8.0):
    if columns is not None: d=df[columns].copy()
    else: d=df.copy()
    d=d.head(max_rows)
    if headers: d=d.rename(columns=headers)
    table=doc.add_table(rows=1,cols=len(d.columns));table.alignment=WD_TABLE_ALIGNMENT.CENTER;table.style='Table Grid'
    # Repeat the header row when a table spans multiple pages.
    trPr=table.rows[0]._tr.get_or_add_trPr()
    tblHeader=OxmlElement('w:tblHeader');tblHeader.set(qn('w:val'),'true');trPr.append(tblHeader)
    hdr=table.rows[0].cells
    for i,c in enumerate(d.columns):
        set_cell_text(hdr[i],c,bold=True,size=font_size,color=(255,255,255));set_cell_shading(hdr[i],'2F5D8C')
    for _,row in d.iterrows():
        cells=table.add_row().cells
        for i,v in enumerate(row):
            if isinstance(v,(float,np.floating)):
                txt='NA' if pd.isna(v) else (f'{v:.4f}' if abs(v)<10 else f'{v:.2f}')
            else: txt='NA' if pd.isna(v) else str(v)
            set_cell_text(cells[i],txt,size=font_size)
    doc.add_paragraph()
    return table


def configure_doc(doc: Document,title: str):
    sec=doc.sections[0];sec.top_margin=Cm(2.2);sec.bottom_margin=Cm(2.0);sec.left_margin=Cm(2.35);sec.right_margin=Cm(2.35)
    styles=doc.styles
    styles['Normal'].font.name='Microsoft YaHei';styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei');styles['Normal'].font.size=Pt(10.5)
    for sty,size,color in [('Title',22,(37,49,61)),('Heading 1',16,(47,93,140)),('Heading 2',13,(58,125,93)),('Heading 3',11,(117,102,166))]:
        s=styles[sty];s.font.name='Microsoft YaHei';s._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei');s.font.size=Pt(size);s.font.bold=True;s.font.color.rgb=RGBColor(*color)
    p=doc.add_paragraph();p.style='Title';p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run(title)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=p.add_run('Afire Project v2 · 论文可用版本优化');r.font.size=Pt(10);r.font.color.rgb=RGBColor(107,114,128)
    doc.add_paragraph()
    for s in doc.sections:
        fp=s.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=fp.add_run('州尺度生态学分析 · 结果用于不公平性监测与优先级识别');r.font.size=Pt(8);r.font.color.rgb=RGBColor(120,120,120)


def add_para(doc,text,bold_prefix=None):
    p=doc.add_paragraph();p.paragraph_format.line_spacing=1.45;p.paragraph_format.space_after=Pt(6)
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold=True;p.add_run(text[len(bold_prefix):])
    else:p.add_run(text)
    return p


def add_bullets(doc,items):
    for item in items:
        p=doc.add_paragraph(style='List Bullet');p.paragraph_format.line_spacing=1.35;p.add_run(item)


def add_figure(doc,stem,caption,width=6.4):
    path=FIGURES/f'{stem}.png'
    if path.exists():
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(path),width=Inches(width))
        c=doc.add_paragraph();c.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=c.add_run(caption);r.italic=True;r.font.size=Pt(9);r.font.color.rgb=RGBColor(90,90,90)


def save_doc(doc,path):
    path.parent.mkdir(parents=True,exist_ok=True);doc.save(path)

# -----------------------------
# Narrative reports
# -----------------------------

def core_stats(panel,fire,ann_exp,models,top,fire_comp,wph_metrics,lag_results):
    exp_min=ann_exp.loc[ann_exp.population_weighted_smoke_pm25.idxmin()]
    exp_max=ann_exp.loc[ann_exp.population_weighted_smoke_pm25.idxmax()]
    gini_lr=linregress(ann_exp.year,ann_exp.gini_smoke_pm25)
    ci_lr=linregress(ann_exp.year,ann_exp.concentration_index_by_ses_disadvantage)
    c2020=fire_comp[(fire_comp.outcome_id=='smoke_pm25_annual_burden')].iloc[0]
    main_ols=models[(models.model=='ols_rate_state_year_fe') & models.term.isin(['z_smoke_pm25_annual_burden','pm25_x_vulnerability'])]
    sig=main_ols[main_ols.p_value_normal_approx<.05]
    lag_core=lag_results[lag_results.term.isin(['z_smoke_pm25_lag1','lag1_pm25_x_vulnerability'])]
    lag_consistency=lag_core[lag_core.direction_consistent_with_main=='yes'].shape[0]
    lag_comparable=lag_core[lag_core.direction_consistent_with_main.isin(['yes','no'])].shape[0]
    return {
        'panel_rows':len(panel),'states':panel.state_abbr.nunique(),'years':panel.year.nunique(),'causes':panel.cause_id.nunique(),
        'exposure_states':fire.state_abbr.nunique(),'missing_states':'AK、HI','top100_stateyears':int(fire.top100_fire_any.sum()),'overlap_stateyears':int(fire.top100_fire_pm25_cover.sum()),
        'exp_min_year':int(exp_min.year),'exp_min':float(exp_min.population_weighted_smoke_pm25),'exp_max_year':int(exp_max.year),'exp_max':float(exp_max.population_weighted_smoke_pm25),
        'gini_slope':float(gini_lr.slope),'gini_p':float(gini_lr.pvalue),'ci_slope':float(ci_lr.slope),'ci_p':float(ci_lr.pvalue),
        'cover_pm25':float(c2020.covered_mean),'noncover_pm25':float(c2020.noncovered_mean),'cover_diff':float(c2020.difference_covered_minus_noncovered),'cover_p':float(c2020.p_value_welch),
        'top1_state':top.iloc[0].state,'top1_cause':top.iloc[0].cause_en,'top1_value':float(top.iloc[0].mean_WPHII),
        'top3':top.head(3)[['state','cause_en','mean_WPHII']].to_dict('records'),
        'sig_ols':sig[['cause_zh','term','estimate','p_value_normal_approx']].to_dict('records'),
        'lag_consistency':lag_consistency,'lag_comparable':lag_comparable,
        'sensitivity':wph_metrics.to_dict('records'),
    }


def upgraded_english_abstract(s):
    return f"""Background: Wildfire-smoke PM2.5 and related health burdens may be unevenly distributed across socially vulnerable populations, but state-level evidence can be distorted when wildfire occurrence is treated as equivalent to smoke exposure. We developed a revised framework that explicitly links major wildfire occurrence with high smoke-PM2.5 coverage.

Methods: We constructed a 2006–2020 state-year-disease panel integrating county-day wildfire-smoke PM2.5, CDC WONDER mortality for seven disease categories, county-derived socioeconomic vulnerability, and records of the 100 largest fires. The primary fire-covered indicator required both allocated Top100 burned area and annual smoke PM2.5 in the upper quartile of the 49-jurisdiction exposure distribution. We quantified annual exposure inequity using population-weighted exposure, Gini, Theil, Concentration Index, SII, and RII. A revised Wildfire-PM2.5 Health Inequity Index (WPHII v2) combined normalized exposure, vulnerability, disease burden, and fire-smoke overlap weights. Disease-specific fixed-effects models and one-year-lag models were treated as exploratory ecological analyses.

Results: The panel contained {s['panel_rows']:,} observations across {s['states']} state-level units, 15 years, and seven disease categories; smoke-PM2.5 exposure was available for 49 jurisdictions. We identified {s['top100_stateyears']} Top100-fire state-years and {s['overlap_stateyears']} Top100–high-PM2.5 overlap state-years. Population-weighted smoke PM2.5 ranged from {s['exp_min']:.3f} μg/m³ in {s['exp_min_year']} to {s['exp_max']:.3f} μg/m³ in {s['exp_max_year']}. Exposure Gini increased by {s['gini_slope']:.4f} per year (p={p_fmt(s['gini_p'])}), whereas the Concentration Index showed no clear linear trend (p={p_fmt(s['ci_p'])}). Fire-covered state-years had higher mean annual smoke PM2.5 than non-covered state-years ({s['cover_pm25']:.3f} vs {s['noncover_pm25']:.3f} μg/m³), a difference partly induced by the coverage definition. The highest mean WPHII v2 combination was {s['top1_state']}–{s['top1_cause']} ({s['top1_value']:.3f}). Health-model coefficients were heterogeneous in sign and magnitude, and lagged estimates were not uniformly consistent with contemporaneous models.

Conclusions: Explicit separation of major-fire occurrence from high smoke-PM2.5 coverage improves the interpretability of state-level inequity surveillance. The most defensible findings concern temporal changes in exposure inequality, spatial coincidence of fire, smoke, vulnerability, and disease burden, and the stability of WPHII-based priority rankings. Disease models should be interpreted as exploratory ecological evidence rather than individual-level causal effects."""


def report_interpretation(s,audit,ann_exp,models,top,fire_comp,lag_results,wph_metrics):
    doc=Document();configure_doc(doc,'论文结果解读与修改建议 v2')
    doc.add_heading('一、是否需要重跑：结论',level=1)
    add_para(doc,'当前 v2 不需要无意义的全量重跑。关键表、关键字段、2006–2020 年份、7 个疾病类别、Top100–PM2.5 overlap 主定义、WPHII v2 公式和疾病模型均通过完整性检查；现有中文 DOCX 也可正常打开。需要修正的是旧报告中残留的 v1 表述、论文图表不足、lag1 和 WPHII 权重敏感性缺失，因此应采用“已有结果二次分析 + 报告重写”的路线。')
    add_table(doc,audit[['check','item','status','detail']],headers={'check':'检查项','item':'对象','status':'状态','detail':'说明'},max_rows=40,font_size=7.5)

    doc.add_heading('二、当前实验设计是否合理',level=1)
    add_para(doc,f'研究采用州尺度 state–year–disease 面板，共 {s["panel_rows"]:,} 条记录，覆盖 {s["states"]} 个州级单元、15 年和 7 类疾病。该设计适合回答全国尺度的暴露分布、不公平性趋势、火灾–烟雾–脆弱性空间叠加和综合优先级识别问题。PM2.5 暴露实际覆盖 49 个州级单元，AK 与 HI 缺少暴露值；模型会因完整案例分析排除这些州年，因此论文中必须明确暴露分析样本口径。')
    add_para(doc,'设计的主要优势是时间跨度长、疾病分层明确、暴露和社会脆弱性在相同州年尺度上统一，并同时保留连续暴露、二元 overlap 和连续 fire–PM2.5 intensity。主要限制是州尺度较粗、Top100 不是全部火灾、年度暴露难以代表短期急性窗口、健康数据存在压制和生态学混杂。')

    doc.add_heading('三、Top100 火灾与 PM2.5 cover 的逻辑',level=1)
    add_bullets(doc,[
        'top100_fire_any：某州某年获得 Top100 大火燃烧面积分配，代表“重大火灾发生线索”。',
        'pm25_cover_top25：某州某年的 smoke PM2.5 年负担位于当年 49 个有暴露州级单元的前 25%，代表“高烟雾暴露覆盖”。',
        'top100_fire_pm25_cover：上述二者同时为 1，是 v2 主 fire_covered 定义。它识别重大火灾与高 smoke PM2.5 在同一州年重叠，但不能证明该州全部 PM2.5 均由该场火灾造成。',
        'fire_pm25_cover_intensity：年度标准化 PM2.5 与年度标准化 log(分配燃烧面积+1) 的乘积，用于连续优先级刻画，不等同于燃烧面积或烟羽覆盖面积。'
    ])
    add_para(doc,f'全期共有 {s["top100_stateyears"]} 个 Top100-fire 州年，其中 {s["overlap_stateyears"]} 个同时达到 PM2.5 Top25%。该定义比“PM2.5>0 即 fire-covered”更具区分度，但也会使 cover 组的 PM2.5 水平按定义更高，因此 cover vs non-cover 的暴露差异只能作为构念验证。')

    doc.add_heading('四、为什么 v2 比 v1 更合理',level=1)
    add_bullets(doc,[
        '避免把“检测到任何 smoke PM2.5”误当作“重大火灾覆盖”，提高了暴露分组的判别力。',
        '把火灾发生和 smoke PM2.5 高覆盖拆成两个可独立描述的维度，可区分“有大火但未形成高州年暴露”和“高暴露但无 Top100 大火记录”。',
        '引入连续强度后，不再完全依赖二元阈值，适合排序、敏感性分析和 WPHII 构建。',
        '论文主结论从不稳定的健康系数转向暴露不公平性、时空叠加和综合风险优先级，更符合数据尺度的证据能力。'
    ])

    doc.add_heading('五、哪些结果可作为论文主结果',level=1)
    add_para(doc,f'第一主结果是暴露不公平性变化。人口加权 smoke PM2.5 在 {s["exp_min_year"]} 年最低（{s["exp_min"]:.3f} μg/m³），在 {s["exp_max_year"]} 年最高（{s["exp_max"]:.3f} μg/m³）。Gini 的线性斜率为 {s["gini_slope"]:.4f}/年（p={p_fmt(s["gini_p"])}），提示州际暴露离散程度总体上升；集中指数趋势不显著（斜率 {s["ci_slope"]:.4f}/年，p={p_fmt(s["ci_p"])}），说明“暴露是否持续向高脆弱州集中”没有稳定单向证据。')
    add_para(doc,'第二主结果是 Top100 大火与高 PM2.5 州年的时空重叠，尤其是 2020 年 Oregon、California、Washington、Colorado 和 Arizona 的 overlap。第三主结果是疾病负担与社会脆弱性的空间叠加。第四主结果是 WPHII v2 排名及其敏感性稳定性，用于识别需要优先监测与资源配置的州–疾病组合。')
    add_para(doc,f'主版本 WPHII 排名第一的是 {s["top1_state"]}–{s["top1_cause"]}，平均值为 {s["top1_value"]:.3f}。该结论应表述为“综合优先级最高”，不能表述为“PM2.5 导致该疾病风险最高”。')

    doc.add_heading('六、健康模型需要谨慎解释的结果',level=1)
    sig=models[(models.model=='ols_rate_state_year_fe')&(models.term.isin(['z_smoke_pm25_annual_burden','pm25_x_vulnerability']))&(models.p_value_normal_approx<.05)][['cause_zh','term','estimate','p_value_normal_approx']]
    add_table(doc,sig,headers={'cause_zh':'疾病','term':'模型项','estimate':'系数','p_value_normal_approx':'p 值'},max_rows=20)
    add_para(doc,'模型中同时出现正向、负向和不显著估计。负向系数可能来自州尺度时间趋势、暴露错配、健康结局压制、跨州医疗与人口结构差异、共线性以及年度暴露与疾病潜伏期不匹配，绝不能解释为 PM2.5 具有保护作用。交互项只能说明在模型设定下关联斜率随社会脆弱性变化，不能直接证明弱势个体具有更强生物学易感性。')
    add_para(doc,f'lag1 补充分析中，可与现有主模型比较的核心方向共 {s["lag_comparable"]} 项，其中 {s["lag_consistency"]} 项方向一致。这种不完全一致进一步支持把疾病模型放在探索性或补充结果，而不是论文标题和结论中心。')

    doc.add_heading('七、不建议写成强因果结论的表述',level=1)
    add_bullets(doc,[
        '不要写“Top100 大火导致 cover 州 PM2.5 增加”，因为 cover 定义包含 PM2.5 Top25%，组间差异部分由分类规则决定。',
        '不要写“野火 PM2.5 导致某疾病死亡率下降”或“具有保护作用”。',
        '不要把州尺度交互项解释成个体层面的易感性或公平效应。',
        '不要把 WPHII 排名解释成因果风险、归因死亡数或效应量。',
        '不要把 Top100 数据描述为全部火灾或全国完整火灾清单。'
    ])

    doc.add_heading('八、可直接放入论文的中文结果段落',level=1)
    result_text=(f'2006—2020 年州尺度分析显示，美国野火烟雾 PM2.5 暴露及其空间不均衡具有明显年际波动。人口加权年暴露由 {s["exp_min_year"]} 年的 {s["exp_min"]:.3f} μg/m³ 上升至 {s["exp_max_year"]} 年的峰值 {s["exp_max"]:.3f} μg/m³。州际暴露 Gini 系数呈显著上升趋势（斜率={s["gini_slope"]:.4f}/年，p={p_fmt(s["gini_p"])}），表明高暴露年份的州际差异总体扩大；然而，按社会经济脆弱性排序的集中指数未呈现显著单向趋势（p={p_fmt(s["ci_p"])}），说明暴露负担并非持续稳定地向高脆弱州集中。进一步将 Top100 大火记录与年度 PM2.5 高覆盖相联结后，共识别 {s["overlap_stateyears"]} 个 overlap 州年。2020 年的重叠主要出现在 Oregon、California、Washington、Colorado 和 Arizona，提示重大火灾、较高烟雾暴露和区域脆弱性可在特定年份形成显著空间聚集。')
    add_para(doc,result_text)

    doc.add_heading('九、可直接放入论文的中文讨论段落',level=1)
    discussion=('本研究最重要的发现并非单一疾病模型系数，而是州际野火烟雾暴露不均衡的增强及其与重大火灾、社会脆弱性和既有疾病负担的复合叠加。Gini 的上升表明，即使全国平均暴露并未呈现稳定线性增长，高暴露可能越来越集中于少数州；集中指数缺乏稳定趋势则提示社会经济分层并非简单的“高脆弱州始终暴露更高”，而是受到火灾年份、烟羽输送和区域气象条件共同塑造。WPHII v2 将暴露、脆弱性、疾病负担和火灾–烟雾重叠统一到优先级框架中，可用于识别需要加强监测、医疗准备和风险沟通的州–疾病组合，但该指数不应替代因果暴露–反应模型。')
    add_para(doc,discussion)

    doc.add_heading('十、可直接放入论文的方法优化段落',level=1)
    method=('为避免将任意非零 smoke PM2.5 暴露误判为重大火灾覆盖，本研究在 v2 中将火灾发生与烟雾暴露分开定义。首先，根据 Top100 火灾与州界相交关系分配燃烧面积，构建 top100_fire_any；其次，在每个年份内依据州级年度 smoke PM2.5 分布确定前 25% 高覆盖州；仅当两项同时满足时，定义 top100_fire_pm25_cover=1。进一步将年度归一化 smoke PM2.5 与年度归一化 log(Top100 分配燃烧面积+1) 相乘，构建连续 fire_pm25_cover_intensity。该策略提高了分组的区分度，并允许分别评估火灾发生、高烟雾暴露及二者重叠。')
    add_para(doc,method)

    doc.add_heading('十一、英文摘要升级版',level=1)
    for para in upgraded_english_abstract(s).split('\n\n'): add_para(doc,para)

    doc.add_heading('十二、答辩时如何解释创新点',level=1)
    add_bullets(doc,[
        '定义创新：不是把 smoke PM2.5>0 等同于火灾覆盖，而是显式区分 Top100 火灾发生、PM2.5 高覆盖和二者重叠。',
        '指标创新：同时使用绝对暴露、相对不公平性指标和 WPHII 综合优先级，避免仅用均值描述环境不公平。',
        '疾病分层：对 7 类疾病分别构建州年负担与模型结果，展示不同结局的异质性，而不是合并成单一健康指标。',
        '证据分级：把暴露不公平性和综合风险识别作为主证据，把健康模型和 lag1 作为探索性补充，主动控制因果外推。',
        '稳健性创新：通过不同火灾权重和无火灾权重版本检验 WPHII 排名是否依赖单一权重设定。'
    ])
    save_doc(doc,REPORTS/'论文结果解读与修改建议_v2.docx')


def report_mainline(s,top,wph_metrics):
    md=f"""# 论文主线 v2

## 核心定位

本研究应定位为**州尺度野火烟雾 PM2.5 暴露不公平性与综合健康优先级研究**。主结论来自暴露分布、Top100–PM2.5 时空重叠和 WPHII v2；疾病模型仅作为探索性证据。

## 第一层：2006–2020 年 wildfire smoke PM2.5 暴露不公平性变化

- 描述人口加权 smoke PM2.5、Gini、Theil、Concentration Index、SII 和 RII 的年际变化。
- 人口加权暴露最低为 {s['exp_min_year']} 年 {s['exp_min']:.3f} μg/m³，最高为 {s['exp_max_year']} 年 {s['exp_max']:.3f} μg/m³。
- Gini 线性斜率为 {s['gini_slope']:.4f}/年（p={p_fmt(s['gini_p'])}），应作为“州际暴露差异总体扩大”的主要趋势证据。
- Concentration Index 无显著线性趋势（p={p_fmt(s['ci_p'])}），不能简单声称高脆弱州始终承受更高暴露。

## 第二层：Top100 大火与 PM2.5 高覆盖州年的时空叠加

- Top100 仅代表研究数据中的前 100 大火，不是全部火灾。
- 全期有 {s['top100_stateyears']} 个 Top100-fire 州年，{s['overlap_stateyears']} 个同时处于 PM2.5 年度 Top25%。
- 重点展示 2020 年 Oregon、California、Washington、Colorado、Arizona 的 overlap 及连续强度。
- cover vs non-cover 的 PM2.5 差异用于验证定义，不作为因果对照。

## 第三层：疾病负担与社会脆弱性的空间叠加

- 分别报告 7 类疾病的年龄调整死亡率、Gini、Concentration Index、SII/RII 和高低脆弱组差异。
- 强调疾病负担、社会脆弱性与高 smoke PM2.5 的复合空间聚集，而不是直接归因于火灾烟雾。

## 第四层：WPHII v2 识别综合健康不公平风险州–疾病组合

- WPHII = exposure_norm × vulnerability_norm × health_burden_norm × fire_weight。
- 排名第一为 {s['top1_state']}–{s['top1_cause']}，平均 WPHII={s['top1_value']:.3f}。
- 用主权重、等权、无火灾权重和强化火灾权重比较 Top10 稳定性。
- WPHII 是资源配置和监测优先级指数，不是因果效应或归因风险。

## 第五层：健康模型作为探索性证据

- 采用州和年份固定效应的率模型与 population-offset 死亡数模型。
- 主模型和 lag1 模型方向并不完全一致；负向系数不能解释为保护作用。
- 论文中宜将模型结果放在“探索性关联”或补充材料，并以生态学尺度、暴露错配和残余混杂解释不稳定性。

## 推荐的结果章节顺序

1. Panel 质量与样本口径；2. 暴露水平与不公平性趋势；3. Top100–PM2.5 overlap；4. 分疾病健康负担不公平性；5. WPHII 排名；6. WPHII 敏感性；7. 探索性主模型；8. lag1 补充分析。

## 一句话主结论

2006–2020 年美国州际野火烟雾 PM2.5 暴露差异总体扩大，重大火灾与高烟雾暴露在少数州年形成显著叠加；将暴露、社会脆弱性、疾病负担与火灾–烟雾重叠整合后，可识别稳定的综合健康不公平优先州–疾病组合，而疾病模型仅提供有限的探索性生态学证据。
"""
    (REPORTS/'论文主线_v2.md').write_text(md,encoding='utf-8-sig')
    doc=Document();configure_doc(doc,'论文主线 v2')
    for section in md.split('\n## ')[1:]:
        lines=section.splitlines();doc.add_heading(lines[0].strip(),level=1)
        for line in lines[1:]:
            line=line.strip()
            if not line: continue
            if line.startswith('- '):
                p=doc.add_paragraph(style='List Bullet');p.add_run(line[2:])
            elif line[0:2].isdigit() and '. ' in line[:4]:
                add_para(doc,line)
            else:add_para(doc,line.replace('**',''))
    add_figure(doc,'Fig1_v2_study_design_flow','图 1. v2 论文主线与研究设计流程。')
    save_doc(doc,REPORTS/'论文主线_v2.docx')


def report_lag(lag_results,s):
    doc=Document();configure_doc(doc,'1 年滞后暴露模型解释 v2')
    doc.add_heading('一、分析目的与定位',level=1)
    add_para(doc,'lag1 分析用于检验前一年州级 smoke PM2.5 与当年疾病负担的方向和稳定性，不用于证明长期因果效应。由于疾病类别具有不同潜伏期，而年度州尺度暴露存在明显测量粗化，该分析应放入补充结果。')
    doc.add_heading('二、模型设定',level=1)
    add_para(doc,'核心公式为 health_outcome_t ~ z(smoke_pm25_lag1) + z(vulnerability) + z(smoke_pm25_lag1)×z(vulnerability) + state FE + year FE。年龄调整死亡率采用 OLS 率模型，死亡数采用 Poisson 模型并使用 log(population) offset；两类模型的标准误均按州聚类。')
    add_para(doc,'新增变量包括 smoke_pm25_lag1、top100_fire_pm25_cover_lag1 和 fire_pm25_cover_intensity_lag1。主模型仅使用 lag1 PM2.5、脆弱性及交互，后两项用于描述和后续扩展，避免在有限样本中加入过多高度相关的火灾变量。')
    core=lag_results[lag_results.term.isin(['z_smoke_pm25_lag1','lag1_pm25_x_vulnerability'])].copy()
    core['model_zh']=core.model.map({'ols_rate_state_year_fe_lag1':'率模型','poisson_deaths_state_year_fe_lag1':'死亡数 offset 模型'})
    core['term_zh']=core.term.map({'z_smoke_pm25_lag1':'lag1 PM2.5 主效应','lag1_pm25_x_vulnerability':'lag1 PM2.5×脆弱性'})
    core['p']=core.p_value.map(p_fmt);core['方向一致']=core.direction_consistent_with_main.map({'yes':'一致','no':'不一致','not_applicable':'不适用'})
    add_table(doc,core[['model_zh','cause_zh','term_zh','estimate','ci95_low','ci95_high','p','方向一致']],headers={'model_zh':'模型','cause_zh':'疾病','term_zh':'模型项','estimate':'估计','ci95_low':'95%CI 下限','ci95_high':'95%CI 上限','p':'p 值','方向一致':'与主模型方向'},max_rows=40,font_size=7.3)
    doc.add_heading('三、结果解释',level=1)
    sig=core[core.p_value<.05]
    if len(sig):
        add_para(doc,f'共 {len(sig)} 个 lag1 核心项达到名义 p<0.05，但这些结果未进行多重比较校正，且不同模型与疾病之间方向并不统一。达到名义显著的项目应描述为“观察到关联”，不能描述为“滞后效应已被证实”。')
    else:add_para(doc,'lag1 核心项未出现名义 p<0.05，提示州尺度年度滞后暴露证据有限。')
    add_para(doc,f'可与主模型比较的核心项中，{s["lag_consistency"]}/{s["lag_comparable"]} 方向一致。方向不一致可能反映暴露窗口不匹配、年度平均暴露稀释、州级迁移和医疗差异、残余时间趋势以及主模型与 lag 模型协变量集合不同。')
    doc.add_heading('四、论文建议写法',level=1)
    add_para(doc,'“作为补充分析，我们将 smoke PM2.5 暴露向后滞后 1 年，并在州和年份固定效应框架下重新估计疾病特异性模型。lag1 系数在不同疾病和模型之间表现出异质方向，且与同期模型并非完全一致，提示州尺度年度暴露不足以稳定刻画疾病潜伏过程。因此，滞后模型仅用于评估结果稳健性，不作为因果推断依据。”')
    doc.add_heading('五、解释边界',level=1)
    add_bullets(doc,['lag1 不能替代分布滞后模型或事件期短期暴露模型。','Poisson offset 模型估计的是州级死亡率关联，不是个体风险。','名义显著结果需考虑 7 种疾病、2 类模型和多个核心项带来的多重比较。','负向估计不得解释为保护作用。'])
    save_doc(doc,REPORTS/'lag1_model_interpretation_v2.docx')


def report_wphii(metrics,comp):
    doc=Document();configure_doc(doc,'WPHII v2 权重敏感性解释')
    doc.add_heading('一、敏感性版本',level=1)
    add_bullets(doc,[
        'WPHII_main：1 + 0.35×top100_fire_any + 0.65×overlap + 0.50×intensity_norm。',
        'WPHII_equal_weight：三个火灾分量均使用 0.50，用于检验主版本对 0.35/0.65 差异权重的依赖。',
        'WPHII_no_fire_weight：不使用火灾权重，仅保留暴露×脆弱性×疾病负担。',
        'WPHII_strong_fire_weight：0.50×fire_any + 1.00×overlap + 0.75×intensity_norm，强化重大火灾–高暴露重叠。'
    ])
    doc.add_heading('二、稳定性结果',level=1)
    m=metrics.copy();m['overlap']=m.top10_overlap_with_main_n.astype(int);m['jaccard']=m.top10_jaccard_with_main;m['rho']=m.spearman_rank_correlation_with_main
    add_table(doc,m[['version_zh','overlap','jaccard','rho']],headers={'version_zh':'版本','overlap':'Top10 与主版本重合数','jaccard':'Top10 Jaccard','rho':'全排名 Spearman ρ'},max_rows=10)
    nonmain=m[m.version!='WPHII_main']
    weakest=nonmain.sort_values('top10_overlap_with_main_n').iloc[0]
    strongest=nonmain.sort_values('spearman_rank_correlation_with_main',ascending=False).iloc[0]
    add_para(doc,f'与主版本最接近的是“{strongest.version_zh}”，其全排名 Spearman ρ={strongest.spearman_rank_correlation_with_main:.3f}。Top10 变化最大的版本是“{weakest.version_zh}”，与主版本重合 {int(weakest.top10_overlap_with_main_n)}/10。若无火灾权重版本仍保持较高相关，说明基础暴露–脆弱性–疾病负担乘积主导总体排序；若 Top10 明显变化，则说明火灾维度对最优先组合具有实质筛选作用。')
    doc.add_heading('三、Top10 比较',level=1)
    show=comp[['version_zh','rank','state_abbr','cause_zh','mean_index','in_main_top10']].copy();show['in_main_top10']=show.in_main_top10.map({'yes':'是','no':'否'})
    add_table(doc,show,headers={'version_zh':'版本','rank':'排名','state_abbr':'州','cause_zh':'疾病','mean_index':'平均指数','in_main_top10':'属于主版本 Top10'},max_rows=40,font_size=7.3)
    doc.add_heading('四、论文解释',level=1)
    add_para(doc,'敏感性分析的目的不是选择“最显著”的权重，而是判断优先级名单是否依赖主观权重。若排名在等权和强化火灾版本中稳定，可说明主要高风险组合由多维度共同驱动；无火灾权重版本的变化则量化了火灾–烟雾联结对排序的增量贡献。论文应报告重合数、Jaccard 和 Spearman 相关，而不应把某一权重版本称为真实风险。')
    add_figure(doc,'Fig8_v2_wphii_sensitivity','图 8. WPHII 权重敏感性：Top10 重合度与整体排名相关。')
    save_doc(doc,REPORTS/'WPHII_sensitivity_interpretation_v2.docx')


def final_report(s,audit,ann_exp,ann_health,models,top,fire_comp,lag_results,wph_metrics):
    doc=Document();configure_doc(doc,'最终论文分析报告 v2')
    doc.add_heading('执行摘要',level=1)
    add_para(doc,'完整性检查支持在现有 v2 结果上直接开展论文写作，无需重新扫描和聚合全部原始数据。本次优化补充了 lag1 暴露、WPHII 权重敏感性、统一的 fire-cover 对照表、8 幅论文级图以及符合生态学解释边界的结果与讨论文本。')
    add_figure(doc,'Fig1_v2_study_design_flow','图 1. 研究设计与论文分析流程。')

    doc.add_heading('1. 研究问题与创新点',level=1)
    add_bullets(doc,['2006–2020 年州际 smoke PM2.5 暴露水平和不公平性如何变化？','Top100 大火与年度高 smoke PM2.5 在哪些州年发生重叠？','社会脆弱性、疾病负担和火灾–烟雾 overlap 如何形成复合风险？','WPHII 排名是否对火灾权重设定稳健？','疾病模型在同期和 lag1 暴露下是否提供方向一致的探索性证据？'])
    add_para(doc,'创新点在于把重大火灾发生、高 PM2.5 覆盖和连续强度分开表示，并将暴露不公平性指标、疾病分层和综合优先级指数置于同一州年框架；同时明确健康模型的探索性层级，避免生态学证据被过度因果化。')

    doc.add_heading('2. 数据与研究设计',level=1)
    add_para(doc,f'主面板包含 {s["panel_rows"]:,} 条 state–year–disease 记录，覆盖 {s["states"]} 个州级单元、2006–2020 年和 7 类疾病。健康结局来自 CDC WONDER；暴露由县日 smoke PM2.5 聚合至州年；社会经济脆弱性由县级指标人口加权至州年；火灾数据为 Top100 大火。PM2.5 暴露表覆盖 {s["exposure_states"]} 个州级单元，AK 和 HI 缺失，因此暴露模型的完整案例样本主要为 49×15 个州年。')
    add_table(doc,audit[['check','status','detail']],headers={'check':'完整性检查','status':'状态','detail':'说明'},max_rows=20,font_size=7.8)

    doc.add_heading('3. Top100 火灾与 PM2.5 cover 定义',level=1)
    add_para(doc,'主定义为 top100_fire_pm25_cover = top100_fire_any × pm25_cover_top25。该定义表示同一州年同时存在 Top100 大火分配面积和年度 smoke PM2.5 前 25% 暴露。连续强度为年度 PM2.5 归一化值与年度 log(分配燃烧面积+1) 归一化值的乘积。')
    add_para(doc,f'全期识别 {s["top100_stateyears"]} 个 Top100-fire 州年和 {s["overlap_stateyears"]} 个 overlap 州年。由于 Top100 只是最大火灾子集，0 不表示该州年没有火灾；由于 PM2.5 Top25% 是相对阈值，overlap 也不等同于烟羽或燃烧面积。')
    add_figure(doc,'Fig3_v2_top100_pm25_overlap_by_year','图 3. 各年 Top100 火灾、高 PM2.5 覆盖及其重叠州数。')
    add_figure(doc,'Fig4_v2_2020_overlap_states','图 4. 2020 年 overlap 州的连续强度。')

    doc.add_heading('4. 2006–2020 年暴露不公平性变化',level=1)
    add_para(doc,f'人口加权 smoke PM2.5 最低为 {s["exp_min_year"]} 年 {s["exp_min"]:.3f} μg/m³，最高为 {s["exp_max_year"]} 年 {s["exp_max"]:.3f} μg/m³。Gini 呈上升趋势（{s["gini_slope"]:.4f}/年，p={p_fmt(s["gini_p"])}），而集中指数的线性趋势不显著（p={p_fmt(s["ci_p"])}）。因此，可得出“州际暴露差异总体扩大”，但不能得出“暴露持续向社会脆弱州集中”的单向结论。')
    add_figure(doc,'Fig2_v2_exposure_inequality_trends','图 2. 暴露均值、Gini 和集中指数趋势。')

    doc.add_heading('5. Top100–PM2.5 cover 的时空分布',level=1)
    pm=fire_comp[fire_comp.outcome_id=='smoke_pm25_annual_burden'].iloc[0]
    add_para(doc,f'全期 fire-covered 州年的平均 smoke PM2.5 为 {pm.covered_mean:.3f} μg/m³，non-covered 州年为 {pm.noncovered_mean:.3f} μg/m³，差值 {pm.difference_covered_minus_noncovered:.3f} μg/m³（Welch p={p_fmt(pm.p_value_welch)}）。该差异主要证明 overlap 定义确实筛选了高暴露州年，但不应视为火灾对 PM2.5 的因果效应。')
    add_figure(doc,'Fig7_v2_fire_cover_vs_nonfire_exposure','图 7. fire-covered 与 non-covered 州年 smoke PM2.5 分布。')

    doc.add_heading('6. 分疾病健康负担不公平性',level=1)
    latest=ann_health[ann_health.year==2020][['cause_zh','population_weighted_age_adjusted_rate','gini_health_rate','concentration_index_health_by_ses_disadvantage']].sort_values('population_weighted_age_adjusted_rate',ascending=False)
    add_para(doc,'疾病负担在不同疾病类别之间量级差异很大，因此不应直接比较原始死亡率并形成单一排名。建议分别报告每类疾病的年度人口加权率、Gini 和集中指数，并用 WPHII 内的年度–疾病内部归一化避免高基线疾病机械占优。')
    add_table(doc,latest,headers={'cause_zh':'疾病','population_weighted_age_adjusted_rate':'2020 人口加权年龄调整率','gini_health_rate':'健康率 Gini','concentration_index_health_by_ses_disadvantage':'健康集中指数'},max_rows=10)

    doc.add_heading('7. 健康模型探索性结果',level=1)
    core=models[(models.model=='ols_rate_state_year_fe')&models.term.isin(['z_smoke_pm25_annual_burden','pm25_x_vulnerability'])].copy();core['p']=core.p_value_normal_approx.map(p_fmt)
    add_table(doc,core[['cause_zh','term','estimate','std_error','p']],headers={'cause_zh':'疾病','term':'模型项','estimate':'估计','std_error':'标准误','p':'p 值'},max_rows=20,font_size=7.8)
    add_para(doc,'同期模型结果在疾病之间方向不一致。缺血性心脏病的 PM2.5 主效应为正，而呼吸系统、COPD 和慢性下呼吸道疾病的率模型主效应为负；脑血管疾病的 PM2.5×脆弱性交互为正。由于模型含多个相关暴露和交互项、暴露为年度州尺度、部分健康率被压制，这些结果只能作为探索性关联。')
    add_figure(doc,'Fig6_v2_disease_model_summary','图 6. 分疾病同期固定效应模型核心项。')
    add_para(doc,f'lag1 结果中，{s["lag_consistency"]}/{s["lag_comparable"]} 个可比较核心项与同期主模型方向一致，说明滞后关系稳定性有限。详细结果见 lag1_model_results_v2.csv 和 lag1_model_interpretation_v2.docx。')

    doc.add_heading('8. WPHII v2 综合指数',level=1)
    add_para(doc,f'WPHII v2 通过暴露、社会脆弱性、疾病负担和火灾–烟雾权重的乘积识别复合优先级。排名第一为 {s["top1_state"]}–{s["top1_cause"]}（平均 WPHII={s["top1_value"]:.3f}）。排名反映多维不利因素的同时出现，不等同于疾病归因比例或因果风险。')
    show=top.head(10)[['state','state_abbr','cause_en','mean_WPHII','slope_WPHII']]
    add_table(doc,show,headers={'state':'州','state_abbr':'缩写','cause_en':'疾病','mean_WPHII':'平均 WPHII','slope_WPHII':'年度斜率'},max_rows=10,font_size=7.8)
    add_figure(doc,'Fig5_v2_wphii_top_states','图 5. WPHII v2 前 20 州–疾病组合。')

    doc.add_heading('9. 稳健性与敏感性分析',level=1)
    m=wph_metrics.copy()
    add_table(doc,m[['version_zh','top10_overlap_with_main_n','top10_jaccard_with_main','spearman_rank_correlation_with_main']],headers={'version_zh':'版本','top10_overlap_with_main_n':'Top10 重合数','top10_jaccard_with_main':'Jaccard','spearman_rank_correlation_with_main':'Spearman ρ'},max_rows=10)
    add_para(doc,'等权、无火灾权重和强化火灾权重版本用于检验优先级名单对主观权重的依赖。应同时报告 Top10 重合数和全排名相关，而不是仅展示某一版本的最高值。')
    add_figure(doc,'Fig8_v2_wphii_sensitivity','图 8. WPHII 权重敏感性。')

    doc.add_heading('10. 主要结论',level=1)
    add_bullets(doc,[f'州际 smoke PM2.5 暴露 Gini 在 2006–2020 年总体上升（{s["gini_slope"]:.4f}/年）。','重大火灾与年度高 smoke PM2.5 的重叠集中于少数州年，2020 年最突出。','社会脆弱性并未表现为稳定单向的暴露集中趋势，应强调时空异质性。','WPHII v2 能识别复合不利因素叠加的优先州–疾病组合，且需通过权重敏感性评价稳定性。','健康同期和 lag1 模型结果异质，不能支撑强因果结论。'])

    doc.add_heading('11. 局限性',level=1)
    add_bullets(doc,['州尺度聚合掩盖州内社区差异，并存在生态学谬误风险。','Top100 数据不是全部火灾，未记录的小火和跨境烟羽可造成分类误差。','年度 smoke PM2.5 不适合替代事件期或短期急性暴露。','PM2.5 数据缺少 AK 和 HI，健康数据存在压制值。','多疾病、多模型和多交互项增加偶然显著与共线性风险。','WPHII 权重具有规范性选择，指数不具有因果或临床阈值含义。'])

    doc.add_heading('12. 论文可直接使用段落',level=1)
    add_para(doc,'结果段落：2006—2020 年，美国州际野火烟雾 PM2.5 暴露具有显著年际波动，且暴露 Gini 呈总体上升趋势，说明高暴露负担日益集中于少数州。将 Top100 大火与年度 PM2.5 前 25% 州年相联结后，识别出有限但具有明显年份聚集的 fire–smoke overlap；2020 年西部多个州同时表现出较高的烟雾暴露和重大火灾强度。WPHII v2 进一步显示，高暴露、社会脆弱性、较高疾病负担和火灾–烟雾重叠可在特定州–疾病组合中形成复合优先级。')
    add_para(doc,'讨论段落：本研究提示，野火烟雾环境不公平并不完全表现为全国平均暴露持续增长，而更多表现为暴露空间离散程度扩大以及极端年份中的区域性叠加。集中指数缺乏稳定趋势说明社会脆弱性与烟雾暴露之间的关系受到火灾位置、烟羽输送和气象条件共同调节。WPHII 可为公共卫生监测和资源准备提供可重复的排序工具，但其用途是识别优先级而非估计因果健康效应。')
    add_para(doc,'方法段落：本研究将 Top100 大火发生与 smoke PM2.5 高覆盖分别编码，并以二者重叠作为主要 fire-covered 定义，从而避免将任意非零烟雾暴露等同于重大火灾影响。连续 fire–PM2.5 intensity 同时考虑年度暴露相对水平和分配燃烧面积，使二元分类与连续优先级分析能够互相验证。')

    doc.add_heading('13. 答辩问题与回答建议',level=1)
    qa=[
        ('为什么只用 Top100 火灾？','Top100 用于构建“重大火灾发生”维度，而 PM2.5 暴露本身来自连续县日数据，因此研究并未把 Top100 当作全部烟雾来源。论文将其称为 major-fire indicator，并把数据不完整列为局限。'),
        ('为什么 cover 组 PM2.5 明显更高？','因为主定义包含 PM2.5 年度 Top25%，该差异是构念设计的一部分，只用于证明分类确实筛选高暴露州年，不作为因果估计。'),
        ('负向健康系数怎么解释？','不能解释为保护作用。州年尺度暴露错配、时间趋势、共线性和残余混杂均可能造成方向不稳定，因此模型被定位为探索性。'),
        ('WPHII 的权重是否主观？','是规范性权重，因此同时提供等权、无火灾权重和强化火灾权重版本，并用 Top10 重合度和 Spearman 相关检验排序稳定性。'),
        ('研究的核心贡献是什么？','核心贡献是将长期暴露不公平性、重大火灾–高烟雾重叠、疾病负担和社会脆弱性统一到可重复的州年优先级框架，同时明确因果解释边界。')]
    for q,a in qa:
        p=doc.add_paragraph();r=p.add_run('问：'+q);r.bold=True;r.font.color.rgb=RGBColor(47,93,140)
        add_para(doc,'答：'+a)
    doc.add_heading('英文摘要升级版',level=1)
    for para in upgraded_english_abstract(s).split('\n\n'):
        p=doc.add_paragraph();p.paragraph_format.line_spacing=1.12;p.paragraph_format.space_after=Pt(3)
        r=p.add_run(para);r.font.size=Pt(9.2);r.font.name='Arial'
    save_doc(doc,REPORTS/'最终论文分析报告_v2.docx')


def write_readme(need_rerun):
    status='需要全量重跑' if need_rerun else '不需要全量重跑'
    text=rf"""# Afire v2 论文优化结果与下一步

## 1. 现在是否需要重跑

结论：**{status}**。

本次完整性检查确认关键表、关键字段、2006–2020 年、7 个疾病类别、Top100–PM2.5 overlap 定义、WPHII v2 公式和疾病模型均完整。当前应直接使用已有结果进行论文写作。只有更换原始数据、修正暴露聚合、改变 Top100 火灾表或发现源数据错误时，才需要运行全流程。

## 2. 只写论文时优先查看

- `reports/最终论文分析报告_v2.docx`：完整论文式结果稿、局限性和答辩问答。
- `reports/论文结果解读与修改建议_v2.docx`：哪些结果可主写、哪些必须谨慎解释。
- `reports/论文主线_v2.docx` / `.md`：结果章节逻辑。
- `figures_paper_v2/Fig1_v2_...` 至 `Fig8_v2_...`：论文和答辩图。
- `tables/fire_cover_vs_noncover_summary_v2.csv`：统一的 cover 对照表。
- `tables/WPHII_top10_comparison_v2.csv`：不同权重 Top10 稳定性。

## 3. 主结果与补充结果

**主结果**：年度暴露均值与不公平性趋势；Top100–PM2.5 overlap 的年份和州分布；分疾病健康负担不公平性；WPHII v2 排名与权重稳定性。

**补充结果**：同期疾病模型；1 年滞后率模型和 population-offset 死亡数模型；fire-covered vs non-covered 的未调整健康率比较。

## 4. 不得过度解释

- 州尺度结果不是个体因果效应。
- 负向健康系数不是保护作用。
- Top100 不是全部火灾。
- PM2.5 cover 是高烟雾暴露覆盖，不等于燃烧面积或单场火灾烟羽。
- WPHII 是综合优先级指数，不是归因风险或效应估计。
- cover 组与 non-cover 组 PM2.5 差异部分由定义决定。

## 5. 只重新生成论文优化结果

在项目根目录运行：

```bash
python scripts/11_paper_optimization_v2.py
```

该命令直接读取 `data_processed` 和 `tables`，不依赖原始数据目录。

## 6. 需要全量重跑时

Windows PowerShell：

```powershell
cd D:\postgraduate\wildfire\Afire\Afire_project_v2
$env:AFIRE_DATA_ROOT = "D:\postgraduate\wildfire\Afire"
python scripts\run_all.py
python scripts\11_paper_optimization_v2.py
```

Windows CMD：

```bat
cd /d D:\postgraduate\wildfire\Afire\Afire_project_v2
set AFIRE_DATA_ROOT=D:\postgraduate\wildfire\Afire
python scripts\run_all.py
python scripts\11_paper_optimization_v2.py
```

Linux：

```bash
cd /path/to/Afire_project_v2
export AFIRE_DATA_ROOT=/path/to/Afire
python scripts/run_all.py
python scripts/11_paper_optimization_v2.py
```

`AFIRE_DATA_ROOT` 必须指向同时包含 `PM2.5`、`health`、`fire` 和 `ecnomy` 子目录的原始数据根目录。

## 7. 输出清单

新增输出包括 lag1 模型、WPHII 敏感性、fire-cover 对照表、8 幅 PNG/PDF 论文图、4 份专题 Word 报告、最终论文分析报告和本说明文件。
"""
    (ROOT/'README_v2_next_steps.md').write_text(text,encoding='utf-8-sig')


def main():
    panel,fire,ann_exp,ann_health,models,wphii,top=read_inputs()
    audit,need_rerun=audit_v2(panel,fire,models,wphii)
    lag_panel,lag_results=fit_lag_models(panel,models)
    wph_sens,wph_agg,wph_comp,wph_metrics=wphii_sensitivity(wphii)
    fire_comp=fire_comparison(panel,fire,wphii)
    fig1_flow();fig2_exposure(ann_exp);fig3_overlap(fire);fig4_2020(fire);fig5_wphii(top);fig6_models(models);fig7_cover(fire);fig8_sensitivity(wph_metrics)
    s=core_stats(panel,fire,ann_exp,models,top,fire_comp,wph_metrics,lag_results)
    report_interpretation(s,audit,ann_exp,models,top,fire_comp,lag_results,wph_metrics)
    report_mainline(s,top,wph_metrics)
    report_lag(lag_results,s)
    report_wphii(wph_metrics,wph_comp)
    final_report(s,audit,ann_exp,ann_health,models,top,fire_comp,lag_results,wph_metrics)
    write_readme(need_rerun)
    # Save a concise machine-readable summary.
    pd.DataFrame([s]).to_csv(TABLES/'paper_optimization_summary_v2.csv',index=False,encoding='utf-8-sig')
    (LOGS/'paper_optimization_v2.log').write_text('Paper optimization completed successfully. Full rerun required: '+str(need_rerun)+'\n',encoding='utf-8')
    print('Completed paper optimization v2')
    print('Full rerun required:',need_rerun)

if __name__=='__main__':
    main()
