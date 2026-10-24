"""Build bilingual SI from authored text and frozen values, never fit or sample."""
from environment import ROOT, SUB, AUD, SD, BASE
import csv, hashlib, json, re, shutil, zipfile
from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from PIL import Image
from lxml import etree
from math_format import math_element, inline, MM_NS, W_NS, MATH_RECORD
from equations import EQUATIONS
from table_assets import prepare_tables

def sha(p): return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def writecsv(path, rows):
    path.parent.mkdir(parents=True,exist_ok=True)
    with path.open('w',encoding='utf-8-sig',newline='') as f:
        csv.writer(f).writerows(rows)

def readcsv(path):
    with Path(path).open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))

def prepare_data():
    full=SD/'full_results';dest=SUB/'supplementary_data'
    jobs=[(1,'state_baseline_poverty.csv'),(2,'state_county_bridge.csv'),(3,'fire_catalogue_subset.csv')]
    result=[]
    duplicate_ranks={r['Rank'] for r in readcsv(full/'fire_potential_duplicates.csv')}
    for num,name in jobs:
        rows=readcsv(full/name)
        if num==1:rows=[r for r in rows if r.get('eligible','').lower()=='true']
        if num==3:
            for row in rows:row['potential_duplicate_name_year']='true' if row['Rank'] in duplicate_ranks else 'false'
        fields=list(rows[0]);path=dest/f'Supplementary_Data_{num}.csv'
        writecsv(path,[fields]+[[r[k] for k in fields] for r in rows])
        result.append({'id':num,'path':str(path),'source':str(full/name),'source_sha256':sha(full/name),'rows':len(rows),'fields':fields})
    assert [r['rows'] for r in result]==[49,72,81]
    (AUD/'supplementary_data_map.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    return result

def setup_doc(lang):
    d=Document();s=d.sections[0]
    s.page_width=Mm(210);s.page_height=Mm(297)
    s.top_margin=s.bottom_margin=Mm(16);s.left_margin=s.right_margin=Mm(20)
    for name in ['Normal','Title','Heading 1','Heading 2','Heading 3','Caption']:
        st=d.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
        st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'SimSun')
        st.font.size=Pt(10.5 if name=='Normal' else 16 if name=='Title' else 12 if name=='Heading 1' else 10.5)
        st.paragraph_format.space_after=Pt(5)
        st.paragraph_format.space_before=Pt(8 if 'Heading' in name else 0)
        st.paragraph_format.keep_with_next='Heading' in name
        st.paragraph_format.widow_control=True
    d.styles['Normal'].paragraph_format.line_spacing=1.1
    for border in list(d.styles.element.xpath('.//w:pBdr')):
        border.getparent().remove(border)
    s.header_distance=Mm(7);s.footer_distance=Mm(7)
    p=s.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');p._p.append(fld)
    d.core_properties.author='';d.core_properties.comments=''
    d.core_properties.subject='AFIRE 2009–2020 five-outcome Supplementary Information'
    return d

def para(doc,text,size=None,style=None):
    p=doc.add_paragraph(style=style);inline(p,text,size)
    return p

def table(doc,rows,widths=None):
    n=len(rows[0]);w=widths or ([24,73,73] if n==3 else [170/n]*n)
    assert abs(sum(w)-170)<.01
    t=doc.add_table(rows=1,cols=n);t.autofit=False
    for c,width in zip(t.columns,w):c.width=Mm(width)
    for i,row in enumerate(rows):
        cells=t.rows[0].cells if i==0 else t.add_row().cells
        for j,(cell,value) in enumerate(zip(cells,row)):
            value=str(value);cell.width=Mm(w[j]);cell.vertical_alignment=WD_ALIGN_VERTICAL.CENTER
            pr=cell._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),'55');el.set(qn('w:type'),'dxa');margins.append(el)
            pr.append(margins)
            if i==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EEF3');pr.append(shade)
            p=cell.paragraphs[0];p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.line_spacing=1.03;p.paragraph_format.keep_with_next=False
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER if i==0 or len(value)<18 else WD_ALIGN_PARAGRAPH.LEFT
            inline(p,value,9.0 if n>=6 else 9.5,i==0)
        t.rows[i]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    borders=OxmlElement('w:tblBorders')
    for side in ['top','bottom','insideH']:
        x=OxmlElement('w:'+side);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'CDD3D8');borders.append(x)
    t._tbl.tblPr.append(borders)
    return t

def remap_references(text,mapping):
    def conv(m):return '['+', '.join(str(mapping[int(i)]) for i in re.findall(r'\d+',m[1]))+']'
    return re.sub(r'\[(\d+(?:\s*,\s*\d+)*)\]',conv,text)

def make_markdown(lang,tables,figures,refs):
    src=(ROOT/'scripts'/f'content_{lang}.md').read_text(encoding='utf-8')
    src=src.replace('已发表图表','文中图表')
    # Production guidance places separate-data-file descriptions outside the SI PDF.
    data_heading='补充数据说明' if lang=='CN' else 'Supplementary Data descriptions'
    ref_heading='补充参考文献' if lang=='CN' else 'Supplementary References'
    pattern=r'## '+re.escape(data_heading)+r'\n(.*?)\n(?=## '+re.escape(ref_heading)+r')'
    match=re.search(pattern,src,re.S)
    if match:
        (SUB/f'Supplementary_Data_Legends_{lang}.md').write_text('# '+data_heading+'\n'+match[1]+'\n',encoding='utf-8')
        src=re.sub(pattern,'',src,flags=re.S)
    src=re.sub(r'10\^([−-]?\d+)',lambda m:'$10^{'+m[1].replace('−','-')+'}$',src)
    if lang=='CN':src=src.replace('{{EQ:4}}','补充公式4中的a为死因，s为州，H为历史训练年份集合。\n\n{{EQ:4}}')
    else:src=src.replace('{{EQ:4}}','In Supplementary Equation 4, a denotes cause, s state and H the historical training years.\n\n{{EQ:4}}')
    # Supplementary reference numbering follows actual first citation, not an old manuscript list.
    cited=[]
    for part in re.findall(r'\[(\d+(?:\s*,\s*\d+)*)\]',src):
        for v in re.findall(r'\d+',part):
            if int(v) not in cited:cited.append(int(v))
    assert sorted(cited)==list(range(1,13)),cited
    mapping={old:i+1 for i,old in enumerate(cited)}
    src=remap_references(src,mapping)
    src=re.sub(r'Supplementary Equation (\d+)',r'Supplementary Equation (\1)',src)
    table_md=['## '+('补充表' if lang=='CN' else 'Supplementary Tables')]
    for t in tables:
        prefix='补充表' if lang=='CN' else 'Supplementary Table '
        table_md+=['',f"### {prefix}{t['id']} {t[lang.lower()+'_title']}",'',remap_references(t[lang.lower()+'_note'],mapping),'',f"<!-- TABLE:{t['id']} -->"]
        rows=t[lang.lower()+'_rows']
        table_md+=['| '+' | '.join(str(c).replace('|','/') for c in rows[0])+' |','| '+' | '.join(['---']*len(rows[0]))+' |']
        table_md+=['| '+' | '.join(str(c).replace('|','/') for c in r)+' |' for r in rows[1:]]
    src=src.replace('{{TABLES}}','\n'.join(table_md))
    fig_md=['## '+('补充图' if lang=='CN' else 'Supplementary Figures')]
    for k,f in figures.items():
        prefix='补充图' if lang=='CN' else 'Supplementary Fig. '
        fig_md+=['',f'<!-- FIGURE:{k} -->',f"### {prefix}{k} {f[lang.lower()+'_title']}",'',f"![{prefix}{k}]({f['figure_file']})",'',f[lang.lower()+'_caption']]
    src=src.replace('{{FIGURES}}','\n'.join(fig_md))
    for num,lines in EQUATIONS.items():
        name=f'补充公式（{num}）' if lang=='CN' else f'Supplementary Equation ({num})'
        src=src.replace('{{EQ:'+str(num)+'}}',name+'\n\n```latex\n'+'\n'.join(lines)+'\n```')
    byid={int(r['id']):r for r in refs}
    src=src.replace('{{REFERENCES}}','\n\n'.join(f'[{mapping[i]}] {byid[i]["text"]}' for i in cited))
    assert '{{' not in src
    (SUB/f'Supplementary_Information_{lang}.md').write_text(src,encoding='utf-8')
    return src,mapping

def build_doc(lang,text,tables):
    MATH_RECORD.clear();doc=setup_doc(lang);lines=text.splitlines();i=0;current_table=None;figs=[];pending_page=False
    while i<len(lines):
        line=lines[i].strip();i+=1
        if not line:continue
        if line.startswith('<!-- FIGURE:'):
            pending_page=True;continue
        if line.startswith('<!-- TABLE:'):
            current_table=int(re.search(r'\d+',line)[0]);continue
        if line.startswith('```latex'):
            eq=[]
            while not lines[i].startswith('```'):eq.append(lines[i]);i+=1
            i+=1
            for j,expr in enumerate(eq):
                p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after=Pt(3);p.paragraph_format.keep_together=True
                p.paragraph_format.keep_with_next=j<len(eq)-1;p._p.append(math_element(expr))
            continue
        if line.startswith('|'):
            while i<len(lines) and lines[i].startswith('|'):i+=1
            t=next(t for t in tables if t['id']==current_table)
            table(doc,t[lang.lower()+'_rows'],t.get('widths_mm'))
            current_table=None;continue
        img=re.match(r'!\[([^]]*)\]\(([^)]+)\)',line)
        if img:
            path=(SUB/img[2]).resolve()
            with Image.open(path) as im:w,h=im.size
            width=170;height=width*h/w
            if height>195:width*=195/height;height=195
            p=doc.add_paragraph();p.paragraph_format.space_after=Pt(3);p.paragraph_format.keep_with_next=True
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            pic=p.add_run().add_picture(str(path),width=Mm(width));pic._inline.docPr.set('descr',img[1])
            figs.append({'figure':img[1],'path':str(path),'width_mm':width,'height_mm':height,'source_sha256':sha(path)})
            continue
        if line.startswith('#'):
            level=len(line)-len(line.lstrip('#'));title=line.lstrip('# ')
            if title in ('补充图','Supplementary Figures'):continue # each figure carries its own title
            if title in ('补充参考文献','Supplementary References'):pending_page=True
            p=para(doc,title,style='Title' if level==1 else 'Heading 1' if level==2 else 'Heading 2')
            p.paragraph_format.keep_with_next=True
            if pending_page:p.paragraph_format.page_break_before=True;pending_page=False
            continue
        p=para(doc,line)
        if line.startswith(('补充公式','Supplementary Equation')):p.paragraph_format.keep_with_next=True
        next_line=next((x.strip() for x in lines[i:] if x.strip()),'')
        if next_line.startswith('<!-- TABLE:'):p.paragraph_format.keep_with_next=True
        if re.match(r'^\[\d+\]',line):p.paragraph_format.space_after=Pt(6)
    output=SUB/f'Supplementary_Information_{lang}.docx';doc.save(output)
    with zipfile.ZipFile(output) as z:xml=etree.fromstring(z.read('word/document.xml'))
    ns={'w':W_NS,'m':MM_NS};body=''.join(xml.xpath('//w:t/text()',namespaces=ns))
    assert not re.search(r'\\(?:beta|frac|sum|mu|alpha|rho)|\$|```|\{\{',body)
    assert len(doc.tables)==22 and len(doc.inline_shapes)==7
    assert len(xml.xpath('//m:oMath',namespaces=ns))==len(MATH_RECORD)
    audit={'document':str(output),'sha256':sha(output),'tables':len(doc.tables),'table_rows':[len(t.rows)-1 for t in doc.tables],
           'figures':figs,'native_math_objects':len(MATH_RECORD),'equation_numbers':list(EQUATIONS),'status':'STRUCTURE_PASS_VISUAL_PENDING'}
    (AUD/f'build_{lang}.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
    (AUD/f'math_objects_{lang}.json').write_text(json.dumps(MATH_RECORD,ensure_ascii=False,indent=2),encoding='utf-8')
    print(lang,output,len(doc.paragraphs),'paragraphs',len(doc.tables),'tables',flush=True)
    return audit

def main():
    tables=prepare_tables();assert [t['id'] for t in tables]==list(range(1,23))
    figures=json.loads((ROOT/'scripts/figure_captions.json').read_text(encoding='utf-8'))
    refs=json.loads((ROOT/'scripts/references.json').read_text(encoding='utf-8'))
    if isinstance(refs,dict):refs=refs.get('references',list(refs.values()))
    prepare_data()
    maps=[]
    for lang in ['CN','EN']:
        text,mapping=make_markdown(lang,tables,figures,refs);build_doc(lang,text,tables);maps.append(mapping)
    assert maps[0]==maps[1]
    writecsv(AUD/'reference_numbering.csv',[['source_reference_id','SI_reference_number']]+list(maps[0].items()))
    (SUB/'supplement_equations.tex').write_text('% Frozen implementation. Display equation numbers 1–12.\n\n'+'\n\n'.join('% Equation '+str(n)+'\n'+'\n'.join(v) for n,v in EQUATIONS.items()),encoding='utf-8')
    (AUD/'table_assets.json').write_text(json.dumps(tables,ensure_ascii=False,indent=2),encoding='utf-8')
    print('Build complete. Render and inspect both documents before release.',flush=True)

if __name__=='__main__':main()
