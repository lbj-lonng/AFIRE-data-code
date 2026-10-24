"""Validate deliverables and current source hashes. No model or allocation calls."""
from environment import ROOT, SUB, AUD, BASE
import csv, hashlib, json, re, zipfile
from pathlib import Path
from collections import Counter
from lxml import etree
from docx import Document
from pypdf import PdfReader
from math_format import math_element
from equations import EQUATIONS

def math_tree(node):
    return (node.tag,tuple(sorted(node.attrib.items())),node.text or '',tuple(math_tree(c) for c in node))

def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def read(path):
    with Path(path).open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))
def write(path,rows):
    if not rows:return
    with Path(path).open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)
def nums(s):
    s=str(s).replace('−','-').replace(',','')
    s=s.replace('/m³','m-3')
    return Counter(re.findall(r'(?<![A-Za-z])[-+]?\d+(?:\.\d+)?(?:[eE][-+]?\d+)?',s))

def main():
    checks=[];bi=[];artifacts=[]
    def ck(name,ok,detail='',state=None):checks.append(dict(check=name,status=state or ('PASS' if ok else 'FAIL'),detail=detail))
    tables=json.loads((AUD/'table_assets.json').read_text(encoding='utf-8'))
    cfg=json.loads((ROOT/'scripts/scope_config.yaml').read_text(encoding='utf-8'))
    ck('table_continuity',[t['id'] for t in tables]==list(range(1,23)))
    for t in tables:
        for i,(a,b) in enumerate(zip(t['cn_rows'][1:],t['en_rows'][1:]),1):
            av=' '.join(map(str,a));bv=' '.join(map(str,b))
            if t['id']==2:
                bv=re.sub(r'\bThirty\b','30',bv,flags=re.I)
                bv=re.sub(r'\bfive\b','5',bv,flags=re.I)
                bv=bv.replace('May–October','5—10')
            if t['id']==18:
                # Outcome/design labels are semantic, not numeric estimates.
                # Explicit design mapping is checked separately below.
                av=' '.join(map(str,a[2:]));bv=' '.join(map(str,b[2:]))
                names={'完整':'Full sample','删除2020':'Omit 2020','删除CA':'Omit CA','同时删除':'Omit CA and 2020','仅删除CA×2020':'Omit CA in 2020'}
                ck(f'table18_semantic_labels_row_{i}',a[0]==b[0] and names.get(a[1])==b[1])
            ok=nums(av)==nums(bv)
            bi.append(dict(check='table_numeric_tokens',item=f"Table {t['id']} row {i}",status='PASS' if ok else 'FAIL',detail='Display values, signs, precision and intervals compared across languages'))
    ck('bilingual_table_values',all(x['status']=='PASS' for x in bi))
    ck('health_table_rows',[len(t['cn_rows'])-1 for t in tables if 15<=t['id']<=19]==[5,25,40,25,40])
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
    for lang in ['CN','EN']:
        path=SUB/f'Supplementary_Information_{lang}.docx';d=Document(path)
        md=(SUB/f'Supplementary_Information_{lang}.md').read_text(encoding='utf-8')
        with zipfile.ZipFile(path) as z:xml=etree.fromstring(z.read('word/document.xml'))
        body=''.join(xml.xpath('//w:t/text()',namespaces=ns))
        ck(lang+'_native_tables',len(d.tables)==22)
        ck(lang+'_embedded_figures',len(d.inline_shapes)==7)
        ck(lang+'_no_placeholders_or_raw_tex',not re.search(r'\{\{|```|\\(?:frac|beta|alpha|sum)|\$',body))
        ck(lang+'_no_tracked_changes',not xml.xpath('//w:ins|//w:del|//w:highlight',namespaces=ns))
        ck(lang+'_editable_math',len(xml.xpath('//m:oMath',namespaces=ns))>22)
        for num,expressions in EQUATIONS.items():
            label=f'补充公式（{num}）' if lang=='CN' else f'Supplementary Equation ({num})'
            matches=[i for i,p in enumerate(d.paragraphs) if p.text==label]
            ok=len(matches)==1
            if ok:
                for offset,expr in enumerate(expressions,1):
                    nodes=d.paragraphs[matches[0]+offset]._p.xpath('./m:oMath')
                    ok=ok and len(nodes)==1 and math_tree(nodes[0])==math_tree(math_element(expr))
            ck(f'{lang}_native_equation_{num}_matches_source',ok,'Native OMML structure and tokens compared to the archived LaTeX expression.')
        with zipfile.ZipFile(path) as z:
            media=Counter(hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/'))
        expected_media=Counter(sha(SUB/f'figures/Supplementary_Figure_{n}.png') for n in range(1,8))
        ck(lang+'_embedded_figure_hashes',media==expected_media)
        ck(lang+'_all_tables_repeat_headers',all(t._tbl.xpath('./w:tr[1]/w:trPr/w:tblHeader') for t in d.tables))
        expected=[len(t[lang.lower()+'_rows'])-1 for t in tables]
        ck(lang+'_table_rows_match',expected==[len(t.rows)-1 for t in d.tables],str(expected))
        # Reopen actual DOCX cells, not the builder's input strings.
        actual=[[[''.join(c._tc.xpath('.//w:t/text()|.//m:t/text()')) for c in r.cells] for r in t.rows] for t in d.tables]
        for i,t in enumerate(tables):
            for j,row in enumerate(t[lang.lower()+'_rows'][1:],1):
                ck(f'{lang}_saved_table_{i+1}_row_{j}',nums(' '.join(map(str,row)))==nums(' '.join(actual[i][j])))
        if lang=='EN':
            ck('EN_no_untranslated_CJK',not re.search(r'[\u3400-\u9fff]',body))
            for kind,maxid in [('Table',22),('Fig.',7),('Equation',12),('Note',6),('Methods',7),('Data',3)]:
                found=[int(x) for x in re.findall(r'Supplementary '+re.escape(kind)+r'\s*\(?(\d+)',md)]
                ck('EN_'+kind+'_reference_bounds',bool(found) and all(1<=x<=maxid for x in found),str(sorted(set(found))))
            headrefs=re.findall(r'^\[(\d+)\]',md,re.M);ck('reference_list_1_to_12',list(map(int,headrefs))==list(range(1,13)))
            before=md.split('## Supplementary References')[0]
            seen=[]
            for group in re.findall(r'\[(\d+(?:,\s*\d+)*)\]',before):
                for num in map(int,re.findall(r'\d+',group)):
                    if num not in seen:seen.append(num)
            ck('references_first_citation_order',seen==list(range(1,13)),str(seen))
        pdf=path.with_suffix('.pdf');render=AUD/'render'/path.stem
        ck(lang+'_PDF_exists',pdf.exists())
        if pdf.exists():
            pages=PdfReader(pdf).pages
            ck(lang+'_render_pages_complete',len(list(render.glob('page-*.png')))==len(pages),str(len(pages)))
            ck(lang+'_no_empty_PDF_pages',all(len((p.extract_text() or '').strip())>3 for p in pages))
            ck(lang+'_PDF_file_limit',pdf.stat().st_size<30*1024*1024,str(pdf.stat().st_size))
            ck(lang+'_render_matches_final_DOCX',json.loads((render/'render_manifest.json').read_text())['docx_sha256']==sha(path))
            ck(lang+'_delivered_PDF_matches_render',sha(pdf)==sha(render/pdf.name))
        for p in [path,pdf]:
            if p.exists():artifacts.append(dict(file=str(p.relative_to(ROOT)),bytes=p.stat().st_size,sha256=sha(p)))
    a=(SUB/'Supplementary_Information_CN.md').read_text(encoding='utf-8')
    b=(SUB/'Supplementary_Information_EN.md').read_text(encoding='utf-8')
    eqA=re.findall(r'```latex\n(.*?)\n```',a,re.S);eqB=re.findall(r'```latex\n(.*?)\n```',b,re.S)
    bi.append(dict(check='native_display_equations',item='Equations 1–12',status='PASS' if eqA==eqB and len(eqA)==12 else 'FAIL',detail='Identical LaTeX expressions converted to editable OMML in both documents'))
    for fact in ['0.678107947271523','3,072','3,062','15,360','2,128','298','92,160','9,127','3,197','506','0.000386','0.004614','39.417','1.888','1.971983','1.960028','1.391503','1.400846','97.5564']:
        bi.append(dict(check='core_fact_in_both_languages',item=fact,status='PASS' if fact in a and fact in b else 'FAIL',detail='Exact displayed core numeric token occurs in both complete texts'))
    # All independent data files exist and retain expected row counts.
    for num,count in [(1,49),(2,72),(3,81)]:ck('data_'+str(num)+'_rows',len(read(SUB/'supplementary_data'/f'Supplementary_Data_{num}.csv'))==count)
    # Scientific outcomes only: explicit excluded labels are checked internally without printing them in submission.
    patterns=[r'alzheimer',r'\bstroke\b',r'阿尔茨海默',r'脑血管',r'脑卒中']
    for lang,txt in [('CN',a),('EN',b)]:ck(lang+'_only_retained_outcomes',not any(re.search(p,txt,re.I) for p in patterns))
    # Freeze the actual latest author identity, separately from documented external editing.
    identity=json.loads((AUD/'maintext_sync_integrity.json').read_text(encoding='utf-8'))
    source=BASE/'submission/Main_Text_CN_2009_2020_five_outcomes_clean.docx'
    expected='d91b9515fc5129c5b1970767d094109961027fbd9d0ce623993395ea0399c55b'
    ck('latest_author_source_stable_since_sync',sha(source)==expected,sha(source))
    # Consolidate provenance, maintaining the externally changed source as an explicit exception.
    source_records={}
    for row in read(AUD/'data_source_manifest.csv'):
        source_records[row['source']]={'path':row['source'],'expected':row['sha256_before'],'role':'current frozen source data'}
    for row in read(AUD/'history_original_files_hash_check.csv'):
        source_records[row['path']]={'path':row['path'],'expected':row['sha256_initial'],'role':'historical or upstream source'}
    for row in read(SUB/'source_data/tables/full_precision_source_index.csv'):
        source_records.setdefault(row['source_file'],{'path':row['source_file'],'expected':row['sha256'],'role':'full-precision table source'})
    for row in read(AUD/'numerical_consistency_check.csv'):
        if row['check'].startswith('source_unchanged:'):
            p=row['check'].split(':',1)[1]
            source_records.setdefault(p,{'path':p,'expected':row['expected'],'role':'original frozen numerical input'})
    hashes=[]
    for row in source_records.values():
        p=Path(row['path']);now=sha(p) if p.is_file() else ''
        state='PASS' if now==row['expected'] else 'EXTERNAL_AUTHOR_CHANGE' if p==source and now==expected else 'FAIL'
        hashes.append(dict(path=str(p),sha256_initial=row['expected'],sha256_current=now,status=state,source_role=row['role'],agent_source_write=False))
    write(AUD/'source_manifest.csv',hashes);write(AUD/'original_files_hash_check.csv',hashes)
    ck('no_unexpected_source_mutation',all(r['status']!='FAIL' for r in hashes),f'{len(hashes)} files; documented external author edit retained')
    ck('maintext_evidence_interpretation',False,'Author removed abstract uncertainty limitation; minimal suggested restoration requires author decision. See author_action_needed.md.',state='PARTIAL')
    ck('upstream_reproducibility_and_sharing',False,'SES construction, some product snapshots, fire extraction/deduplication, county raw responses and publishing permissions remain incomplete.',state='PARTIAL')
    visual_rows=[]
    for stem,review_name in [('Supplementary_Information_CN','CN'),('Supplementary_Information_EN','EN'),('Main_Text_CN_SI_links_synced','maintext')]:
        review=AUD/f'visual_review_{review_name}.csv'
        records=read(review) if review.exists() else []
        page_images=sorted((AUD/'render'/stem).glob('page-*.png'))
        for p in page_images:
            matching=[r for r in records if Path(r.get('page_image',''))==p and r.get('sha256')==sha(p)]
            inspected=len(matching)==1
            review_status=matching[0].get('status','') if inspected else 'NOT_RUN'
            acceptable=inspected and (review_status=='PASS' or (review_name=='maintext' and matching[0].get('si_reference_layout_status')=='PASS'))
            ck(stem+'_'+p.stem+'_current_visual_review',acceptable,review_status,state=None if inspected else 'NOT_RUN')
            visual_rows.append(dict(document=stem,page=p.stem,page_image=str(p),sha256=sha(p),review_status=review_status,scoped_status='PASS' if acceptable else 'NOT_RUN' if not inspected else 'FAIL',review_file=str(review)))
        manifest=AUD/'render'/stem/'render_manifest.json'
        ck(stem+'_final_render_identity',manifest.exists() and json.loads(manifest.read_text())['docx_sha256']==sha(SUB/(stem+'.docx')))
    write(AUD/'visual_review.csv',visual_rows)
    for lang in ['CN','EN']:
        rows=[r for r in visual_rows if r['document']=='Supplementary_Information_'+lang]
        build_path=AUD/f'build_{lang}.json'
        build=json.loads(build_path.read_text(encoding='utf-8'))
        build['status']='STRUCTURE_AND_FINAL_VISUAL_PASS' if rows and all(r['scoped_status']=='PASS' for r in rows) else 'STRUCTURE_PASS_VISUAL_PENDING'
        build['final_visual_pages']=len(rows)
        build_path.write_text(json.dumps(build,ensure_ascii=False,indent=2),encoding='utf-8')
    for filename in ['source_manifest.csv','version_content_migration.csv','claim_evidence_map.csv','terminology_ledger.csv','old_to_new_numbering.csv','maintext_si_crossrefs.csv','numerical_consistency_check.csv','scope_consistency_check.csv','figure_source_data_map.csv','original_files_hash_check.csv','statistical_boundaries.md','unresolved_items.md','author_action_needed.md','maintext_changes.md','journal_format_check.md','visual_review.csv']:
        ck('required_audit_'+filename,(AUD/filename).is_file())
    write(AUD/'bilingual_consistency_check.csv',bi);write(AUD/'supplement_verification.csv',checks)
    artifacts=[dict(file=str(p.relative_to(ROOT)),bytes=p.stat().st_size,sha256=sha(p)) for p in sorted(SUB.rglob('*')) if p.is_file()]
    write(AUD/'artifact_manifest.csv',artifacts)
    summary={'checks':dict(Counter(r['status'] for r in checks)),'bilingual':dict(Counter(r['status'] for r in bi)),
             'source_hashes':dict(Counter(r['status'] for r in hashes)),'overall':'PARTIAL_AUTHOR_AND_SOURCE_ACTIONS','no_refit_or_imputation':True}
    (AUD/'verification_summary.json').write_text(json.dumps(summary,indent=2),encoding='utf-8')
    print(json.dumps(summary,indent=2))
    if any(r['status']=='FAIL' for r in checks+bi):raise SystemExit(1)

if __name__=='__main__':main()
