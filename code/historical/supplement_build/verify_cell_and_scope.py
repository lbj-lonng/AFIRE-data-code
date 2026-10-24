"""Independent per-cell, publication-scope and source-copy verification.

No fitting, sampling, inference or source-data editing. Output is audit only.
"""
from pathlib import Path
from collections import Counter
import csv
import hashlib
import json
import re
import unicodedata
import zipfile
from lxml import etree
from docx import Document
from pypdf import PdfReader

ROOT=Path(__file__).resolve().parents[1]
SUB=ROOT/'submission'
AUD=ROOT/'audit'
CFG=json.loads((ROOT/'scripts/scope_config.yaml').read_text(encoding='utf-8'))
BASE=Path(CFG['base'])
ORIGINAL_CONFIG=BASE/'audit/inputs/health_config.json'
W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
M='http://schemas.openxmlformats.org/officeDocument/2006/math'


def sha(path):
    h=hashlib.sha256()
    with Path(path).open('rb') as f:
        for chunk in iter(lambda:f.read(1048576),b''):h.update(chunk)
    return h.hexdigest()


def rows(path):
    with Path(path).open(encoding='utf-8-sig',newline='') as f:
        return list(csv.DictReader(f))


def save(name,records):
    if not records:return
    with (AUD/name).open('w',encoding='utf-8-sig',newline='') as f:
        out=csv.DictWriter(f,fieldnames=list(records[0]));out.writeheader();out.writerows(records)


def norm_cell(text):
    """Only presentation equivalences; preserve each cell's entire semantic order."""
    text=unicodedata.normalize('NFKC',str(text)).replace('µ','μ')
    text=text.replace('−','-').replace('‐','-').replace('‑','-')
    text=text.replace('\u200b','').replace('\ufeff','')
    # DOCX may render a native inverse-volume exponent instead of slash notation.
    text=re.sub(r'μ\s*g\s*/\s*m\s*3', 'μg m-3', text)
    text=re.sub(r'μ\s*g\s*m\s*[-−]\s*3', 'μg m-3', text)
    text=re.sub(r'\s+','',text)
    return text


def verify_cells():
    assets=json.loads((AUD/'table_assets.json').read_text(encoding='utf-8'))
    results=[]
    for lang in ['CN','EN']:
        path=SUB/f'Supplementary_Information_{lang}.docx'
        doc=Document(path)
        results.append(dict(file=path.name,table=0,row=0,column=0,status='PASS' if len(doc.tables)==len(assets)==22 else 'FAIL',expected='22 tables',actual=str(len(doc.tables)),rule='table count'))
        for table,asset in zip(doc.tables,assets):
            expected=asset[lang.lower()+'_rows']
            results.append(dict(file=path.name,table=asset['id'],row=0,column=0,status='PASS' if len(table.rows)==len(expected) else 'FAIL',expected=str(len(expected)),actual=str(len(table.rows)),rule='row count including header'))
            for i,(actualrow,expectedrow) in enumerate(zip(table.rows,expected)):
                results.append(dict(file=path.name,table=asset['id'],row=i,column=0,status='PASS' if len(actualrow.cells)==len(expectedrow) else 'FAIL',expected=str(len(expectedrow)),actual=str(len(actualrow.cells)),rule='column count'))
                for j,(cell,value) in enumerate(zip(actualrow.cells,expectedrow)):
                    actual=''.join(cell._tc.xpath('.//w:t/text()|.//m:t/text()'))
                    good=norm_cell(actual)==norm_cell(value)
                    results.append(dict(file=path.name,table=asset['id'],row=i,column=j+1,status='PASS' if good else 'FAIL',expected=str(value),actual=actual,rule='same-language whole cell in order; whitespace/minus/unit presentation normalization only'))
    save('independent_cell_comparison.csv',results)
    return results


def references_only_context(line,section):
    return section=='references' or bool(re.match(r'^\s*\[?\d+\]?\s*[.)]?\s+.{0,160}\b(?:doi|https?://|Nature|Science|Lancet|Environ)',line,re.I))


def scan_scope():
    original=json.loads(ORIGINAL_CONFIG.read_text(encoding='utf-8'))
    forbidden=[]
    for outcome in original['outcomes']:
        if outcome['outcome_id'] not in CFG['outcomes']:
            forbidden.extend([outcome['outcome_id'],outcome['outcome_name']])
    # Identity/name terms only, never ICD ranges of retained parent categories.
    compiled=[(word,re.compile(re.escape(word),re.I) if re.search(r'[\u3400-\u9fff]',word) else re.compile(r'(?<![A-Za-z0-9])'+re.escape(word)+r'(?![A-Za-z0-9])',re.I)) for word in forbidden]
    findings=[];coverage=[]
    def inspect(path,location,text,section='',explicit=False):
        for word,pat in compiled:
            if pat.search(str(text)):
                isref=not explicit and references_only_context(str(text),section)
                findings.append(dict(file=str(path.relative_to(SUB)),location=location,term=word,status='REFERENCE_CONTEXT_REVIEW' if isref else 'FAIL',excerpt=str(text)[:500],rule='excluded outcome ID/name; ICD parent ranges not scanned'))
    for path in sorted(SUB.rglob('*')):
        if not path.is_file():continue
        suffix=path.suffix.lower();status='SCANNED';unit_count=0
        if suffix=='.csv':
            data=rows(path)
            for i,row in enumerate(data,2):
                for key,value in row.items():
                    inspect(path,f'row {i} field {key}',value,'references' if str(key).lower() in ['reference','citation','bibliography','article_title'] else '',explicit=str(key).lower() in ['outcome_id','cause_id','outcome_name','child','parent'])
                    if str(key).lower() in ['outcome_id','cause_id','child','parent'] and value and value not in CFG['outcomes']:
                        findings.append(dict(file=str(path.relative_to(SUB)),location=f'row {i} field {key}',term=value,status='FAIL',excerpt=value,rule='explicit outcome ID is outside current whitelist'))
                unit_count+=1
        elif suffix=='.docx':
            with zipfile.ZipFile(path) as z:
                for part in [n for n in z.namelist() if re.fullmatch(r'word/(?:document|header\d+|footer\d+|footnotes|endnotes|comments)\.xml',n)]:
                    xml=etree.fromstring(z.read(part));section=''
                    for i,p in enumerate(xml.findall('.//{'+W+'}p')):
                        text=''.join(p.itertext()) if False else ''.join(p.xpath('.//w:t/text()|.//m:t/text()',namespaces={'w':W,'m':M}))
                        if re.match(r'^(?:Supplementary References|References|补充参考文献|参考文献)$',text.strip()):section='references'
                        inspect(path,f'{part} paragraph {i+1}',text,section);unit_count+=1
        elif suffix=='.pdf':
            section=''
            for i,page in enumerate(PdfReader(path).pages):
                text=page.extract_text() or ''
                if re.search(r'(?:Supplementary References|补充参考文献|参考文献)',text):section='references'
                for line in text.splitlines():inspect(path,f'page {i+1}',line,section)
                unit_count+=1
        elif suffix in {'.md','.txt','.tex','.bib','.json','.geojson','.xml','.svg'}:
            text=path.read_text(encoding='utf-8-sig',errors='replace');section='references' if suffix=='.bib' else ''
            for i,line in enumerate(text.splitlines(),1):
                if re.match(r'^#+\s+(?:Supplementary References|References|补充参考文献|参考文献)\s*$',line):section='references'
                inspect(path,f'line {i}',line,section);unit_count+=1
        else:
            status='NOT_TEXT_CONTAINER'
        coverage.append(dict(file=str(path.relative_to(SUB)),status=status,units=unit_count,sha256=sha(path),scope='binary geometry/image assets not interpreted as disease text' if status!='SCANNED' else 'entire supported text container'))
    if not findings:
        findings=[dict(file='submission/**',location='all supported containers',term='',status='PASS',excerpt='No excluded outcome ID/name or invalid explicit outcome ID found.',rule='dynamic original-config exclusion against current whitelist')]
    save('independent_outcome_scope_findings.csv',findings)
    save('independent_scope_scan_coverage.csv',coverage)
    return findings,coverage


def verify_copies():
    checks=[]
    for row in rows(AUD/'data_source_manifest.csv'):
        if not row['destination']:continue
        dest=Path(row['destination']);source=Path(row['source'])
        if not dest.is_relative_to(SUB):continue
        current=sha(dest) if dest.is_file() else ''
        good=dest.is_file() and current==row['sha256_before']==sha(source)
        checks.append(dict(file=str(dest.relative_to(SUB)),status='PASS' if good else 'FAIL',role='baseline semantic source byte-copy',source=str(source),detail='original manifest, current original and final delivered hashes compared',observed=current))
    for row in rows(SUB/'source_data/tables/full_precision_source_index.csv'):
        src=Path(row['source_file']);dest=SUB/row['delivered_file'];role=row['role']
        good=False;detail='';observed=''
        if role=='full_precision':
            observed=sha(dest) if dest.is_file() else ''
            good=dest.is_file() and observed==sha(src)==row['sha256'];detail='byte-identical full-precision source copy'
        elif role=='full_precision_historical_selection':
            source=rows(src);selected=rows(dest)
            expected={r['model_id']:r for r in source if r['cause_id']=='copd'}
            unique=len({r['model_id'] for r in selected})==len(selected)
            good=unique and len(selected)==(2 if int(row['table'])==11 else 11) and all(r['model_id'] in expected and r==expected[r['model_id']] for r in selected) and sha(src)==row['sha256']
            detail='every field compared as the exact original string by unique historical COPD model_id; original blanks retained';observed=str(len(selected))+' historical rows'
        elif role=='allocation_identity':
            config=json.loads(src.read_text(encoding='utf-8'));selected=rows(dest)
            good=len(selected)==5 and {r['outcome_id'] for r in selected}==set(CFG['outcomes']) and all(str(config['seed_map'][r['outcome_id']])==r['seed'] and str(config['imputations'])==r['n_allocations'] for r in selected) and sha(src)==row['sha256']
            detail='five original seed identifiers and allocation counts compared, no draws produced';observed=str(len(selected))+' retained identities'
        else:
            checks.append(dict(file=str(dest.relative_to(SUB)),status='NOT_RUN',role=role,source=str(src),detail='Role has no implemented independent comparator',observed=''))
            continue
        checks.append(dict(file=str(dest.relative_to(SUB)),status='PASS' if good else 'FAIL',role=role,source=str(src),detail=detail,observed=observed))
    save('independent_source_copy_check.csv',checks)
    return checks


def main():
    cells=verify_cells();scope,coverage=scan_scope();copies=verify_copies()
    summary=dict(cell_checks=dict(Counter(r['status'] for r in cells)),scope_findings=dict(Counter(r['status'] for r in scope)),scope_containers=dict(Counter(r['status'] for r in coverage)),source_copy_checks=dict(Counter(r['status'] for r in copies)),scope_config_sha256=sha(ROOT/'scripts/scope_config.yaml'),table_assets_sha256=sha(AUD/'table_assets.json'),excluded_terms_source=str(ORIGINAL_CONFIG),not_run=['Cross-language prose semantics are covered by the separate scientific content review, not by this comparator.','Newly constructed figure point tables and Supplementary Data row semantics are delegated to figure/data audits; this script verifies manifested byte-copies and selected historical tables.','Rasterized text or shapefile binary attributes are not OCR-scanned; text containers and figure SVGs/PDFs are scanned.','Final native equation mapping and final image-review coverage are checked by the parent verifier.'])
    (AUD/'independent_verification_summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(summary,ensure_ascii=False,indent=2))
    errors=[r for r in cells+scope+copies if r['status']=='FAIL']
    if errors:
        print(json.dumps(errors[:15],ensure_ascii=False,indent=2))
        raise SystemExit(1)


if __name__=='__main__':main()
