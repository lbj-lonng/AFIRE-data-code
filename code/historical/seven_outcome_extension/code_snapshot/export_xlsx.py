"""Export existing CSVs with installed system openpyxl; no statistical fitting."""
import csv, math
from pathlib import Path
from openpyxl import Workbook,load_workbook
O=Path(__file__).resolve().parent
FILES={'all_results':'AFIRE_7outcome_all_results.csv','primary':'primary_7outcomes.csv','marginal':'marginal_associations_7outcomes.csv','common_sample':'common_sample_7outcomes.csv','omissions':'targeted_omissions_7outcomes.csv','LST_pairs':'LST_paired_7outcomes.csv','multiplicity':'primary_multiplicity_7outcomes.csv','diagnostics':'diagnostics_7outcomes.csv','coverage':'analysis_coverage_after.csv'}
def value(v):
    if not v or v.lower()=='nan':return None
    try:
        a=float(v)
        return a if math.isfinite(a) else v
    except ValueError:return v
w=Workbook();w.remove(w.active)
for name,file in FILES.items():
    s=w.create_sheet(name)
    with (O/file).open() as f:
        rows=list(csv.reader(f));s.append(rows[0])
        for row in rows[1:]:s.append([value(v) for v in row])
    s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
    for col in s.columns:s.column_dimensions[col[0].column_letter].width=min(42,max(14,len(str(col[0].value))+2))
w.save(O/'AFIRE_7outcome_results.xlsx')
v=load_workbook(O/'AFIRE_7outcome_results.xlsx',read_only=True,data_only=True)
for name,file in FILES.items():
    with (O/file).open() as f:rows=list(csv.reader(f))
    actual=list(v[name].values);assert len(rows)==len(actual)
    for i,row in enumerate(rows):
        for j,x in enumerate(row):
            want=x if i==0 else value(x);got=actual[i][j]
            assert got==want or isinstance(want,float) and isinstance(got,(float,int)) and math.isclose(want,got,rel_tol=2e-14,abs_tol=1e-300),(name,i,j,want,got)
print('XLSX all sheets cell-by-cell verified')
(O/'logs/xlsx_verification.log').write_text('XLSX all sheets cell-by-cell verified\n')
