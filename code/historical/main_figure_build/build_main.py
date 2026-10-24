"""Revise the latest complete manuscript; preserve native equations and citations."""
from environment import ROOT,SD,SUB,AUD
from copy import deepcopy
import json,re,hashlib
from docx import Document
from docx.shared import Mm,Pt,RGBColor
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn
from scope import CONFIG

def run():
    source=AUD/'inputs/Main_Text_CN_clean.docx';doc=Document(source)
    assert CONFIG['study_start_year']==2009 and CONFIG['study_end_year']==2020 and len(CONFIG['outcomes'])==5
    ps=list(doc.paragraphs);assert len(ps)==97 and len(doc.inline_shapes)==4
    changes=[]
    def setp(i,s):
        p=ps[i];old=p.text;assert not p._p.xpath('.//m:oMath')
        for c in list(p._p):
            if not c.tag.endswith('}pPr'):p._p.remove(c)
        for part in re.split(r'(\[C:[\d,]+\])',s):
            if not part:continue
            r=p.add_run(part[3:-1] if part.startswith('[C:') else part)
            r.font.name='Times New Roman';r.font.size=Pt(9 if p.style.name=='Caption' else 10.5);r.bold=False;r.font.color.rgb=RGBColor(0,0,0)
            r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'SimSun')
            if part.startswith('[C:'):r.font.superscript=True
        changes.append({'paragraph':i,'old':old,'new':s})
    def rep(i,a,b):setp(i,ps[i].text.replace(a,b))
    rep(3,'2006—2020年','2009—2020年')
    rep(3,'五项交互均通过原七项检验的Holm校正','五项交互均通过保留原检验家族的Holm校正')
    setp(9,'本研究以美国本土相连48州及哥伦比亚特区的暴露资料描述时空背景，并在Arizona、California、Colorado、Idaho、Oregon和Washington六州全部256县中展开比较。研究回答两个问题：烟雾高暴露集中于哪些年份和地区，是否更多分布在社会经济弱势县；在同一县域范围内，烟雾—死亡关联是否呈现暴露分布没有反映的社会经济差异。州、县暴露描述及五类心肺死因的年度死亡分析均限定于2009—2020年，共12年。')
    setp(14,'图1｜州尺度烟雾暴露、经济背景与归档火灾分布。a，49个有暴露值州级单位2009—2020年平均年度烟雾PM₂.₅；六州研究范围以深色描边，Alaska和Hawaii缺失而非零，DC单列放大。b，同期49单位等权均值（实线）、中位数（虚线）和第10—90百分位（阴影）；阴影为州际分布范围，不是置信区间，也非全国人口加权暴露。c，固定基线贫困比例的县人口加权州级代理值与同期暴露均值，49个等面积点，橙色标示六州；不含拟合线或检验，真实基线期见方法。d，从既有归档Top100目录按2009—2020年筛得的81条记录，产生89次州—记录归属；六州为55次，涉及53条不同目录记录。橙色为六州；跨州及潜在同年同名记录均保留，不等于独立火灾数，也不是重新生成的该时期全国Top100。原提取版本和全库去重排名尚待确认（补充表S20—S22）。数值见Source Data。')
    setp(16,'县级分析沿用Arizona、California、Colorado、Idaho、Oregon和Washington六州已经形成的死亡、人口、暴露及协变量联结资料，覆盖其全部256县。归档火灾目录中属于2009—2020年的81条记录产生89次州归属，其中55次归属于上述六州，涉及53条不同目录记录（图1d）。这一子集仅提供火灾背景，不是重新生成的该时期全国Top100。跨州记录及同年同名潜在重复均保留，因此这些数目不等于独立火灾数，也不能作为全国代表性或事前选区规则的证据；目录原始提取版本和全库排名去重流程尚待补齐（补充表S21—S22）。')
    setp(17,'县级资料显示，州平均值内部还存在较大的暴露差异。California各县的长期平均浓度为0.179—2.533 μg m⁻³；六州区域人口加权年度浓度则从2019年的0.133 μg m⁻³升至2020年的2.797 μg m⁻³（图2a、b）。高平均浓度并不意味着更大的县际相对不平等：2020年的县等权Gini为0.350，并非研究期间最高（图2c）。异常火季提高了区域暴露水平，但暴露强度与分布不均衡程度并非同步变化。')
    rep(20,'2006—2020年','2009—2020年')
    setp(22,'本批次州级死亡总量已公开，部分县级小计数未报告具体值。2009—2020年五类结局共15,360条县—年—死因记录，其中3,328条（21.67%）受抑制、325条为公开零值。我们利用州总量与既定计数范围保留受抑制记录，使公开计数（包括零）不变，并传播既定分配方案下的计数不确定性（图3a、b；补充S3）。每项健康参考分析覆盖3,072条县—年记录；上述总记录数不能解释为不重复死亡人数。')
    setp(23,'伪抑制检验提供了有限的支持：筛选既有逐单元预测后，2009—2020年五类结局共有2,128个原本公开的10—19计数单元。固定验证经验贝叶斯分配的平均绝对误差为1.97例，低于等权与人口权重分配的2.23和2.87例；层级验证经验贝叶斯为1.96例。固定法名义95%分配区间的经验覆盖率为97.6%（图3c、d）。该检验使用0—19分配支持及既有历史训练信息，不能等同于真实1—9受抑制记录或死亡模型置信区间的校准，也不证明固定法在所有指标上最优（补充表S13）。')
    setp(25,'图3｜县级死亡压制及州总量约束分配的支持证据。a，2009—2020年原始Suppressed记录比例，每县分母60条县—年—死因记录（12年×5类）；地图为256县，不是插补死亡比例。b，五类结局受抑制记录占各自3,072条记录的比例（蓝）与隐藏余量占对应州死亡总量之和的比例（橙）；两列分母及横轴范围不同，均从零起，嵌套死因不可相加。c、d，同期五类结局验证子集中等权、人口权重、固定验证经验贝叶斯（EB）及层级验证EB的平均绝对误差和名义95%分配区间经验覆盖率。四方法使用完全相同的298组、2,128个公开10—19计数单元，每方法100次已存分配，支持范围0—19；d虚线为95%参考，无额外误差条。评价范围筛选不改变原历史训练资料。验证实现不等同于主分析30份1—9分配，覆盖率亦非死亡模型区间校准。完整指标见补充表S13及Source Data。')
    setp(26,'与长期暴露没有沿县级劣势依次升高不同，五类心肺死因的烟雾—死亡关联均随劣势升高而趋于不利。COPD、慢性下呼吸道疾病、呼吸系统疾病、循环系统疾病和缺血性心脏病的交互均通过保留原检验家族的Holm校正；每一劣势标准差对应的死亡率比相对变化为0.299%—0.752%（图4a、b；补充表S5）。五类死因存在包含关系，不是五次独立重复。')
    setp(30,ps[30].text.replace('固定2006—2010年基线综合劣势','固定基线综合劣势（真实基线期见方法）').replace('仅展示五类不改变原七项参考交互的Holm家族；哮喘、脑血管和全部敏感性结果见补充表S5—S9。','估计覆盖2009—2020年；保留原Holm校正值及其真实家族定义，见统计方法。五类结局全部条件关联和敏感性结果见补充表S5—S9。'))
    setp(43,'本研究观察期统一为2009—2020年，共12年。州级暴露描述覆盖美国本土相连48州及哥伦比亚特区，共49个州级单位和588个有暴露观测州—年；Alaska和Hawaii无暴露值。县级研究包括Arizona、California、Colorado、Idaho、Oregon和Washington全部256县，暴露描述、死亡描述及参考分析均使用同期3,072个县—年。县域提供了连接死亡、人口、社会经济与环境资料的共同单元，兼顾州内差异与非城市地区的覆盖；本研究未进行城市、县与其他单元之间的最优尺度比较。')
    setp(44,'火灾背景从既有归档Top100目录按事件年份筛选2009—2020年记录，保留81条及其原始Rank、SizeMetricAcres面积字段和标识，不重排或重新宣称为该时期全国Top100。按已存相交州名单计数，缺少相交州时才使用质心州；跨州记录在每一所属州各计一次。目录仅描述现有研究区域的火灾背景，不据此新增、排除或重新选择研究州。原始提取日期、版本与全库去重流程仍待确认，潜在同名同年记录不擅自合并（补充表S21—S22）。')
    setp(46,'本研究保留五类心肺死因：COPD（J40—J44）、慢性下呼吸道疾病（J40—J47）、呼吸系统疾病（J00—J98）、循环系统疾病（I00—I99）和缺血性心脏病（I20—I25）。这些结局既包含具体病种，也包含较广的心肺死因组，与已有颗粒物及烟雾健康研究关注的结局相衔接[C:5,6,7]。所有代码沿用原下载定义，不从较广类别中删减任何组成诊断；类别彼此嵌套，不能相加为独立死亡（补充表S1）。')
    # Paragraph 58 includes native math: edit only ordinary Word text nodes.
    for node in ps[58]._p.xpath('.//w:t'):
        if node.text:node.text=node.text.replace('Holm校正仍针对原七项参考交互，不因正文聚焦五类而重算','本稿报告五类结局，但保留既有Holm校正值；其原检验家族含7项参考交互，不按当前报告范围重算')
    setp(61,'四张主图和补充图的数值、五类结局完整135行已存健康结果、12年暴露描述及筛选后的验证记录随Source Data提供。上游资料包括CDC WONDER聚合死亡与人口、公开野火烟雾产品及县级社会经济与环境数据。受抑制计数的分配不是隐藏死亡真值，不默认公开逐份分配数组。归档火灾目录的原始提取版本和全库去重排名流程尚未闭合。【待补：公共数据仓库标识、部分精确输入版本、CDC原始查询归档和再分发许可。】')
    for i,n in [(13,1),(19,2),(24,3),(29,4)]:
        p=ps[i];old=doc.inline_shapes[n-1] if False else None
        for el in list(p._p):
            if not el.tag.endswith('}pPr'):p._p.remove(el)
        pic=p.add_run().add_picture(str(SUB/f'figures/Figure_{n}.png'),width=Mm(170))
        pic._inline.docPr.set('descr',f'Figure {n}; 2009–2020; five outcomes where applicable; panels a–d')
    # Keep the final reference with the reference list rather than on a nearly empty page.
    for p in ps[67:]:p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(0)
    used={x.get(qn('r:embed')) for x in doc.element.xpath('.//a:blip')}
    for rid,rel in list(doc.part.rels.items()):
        if rel.reltype.endswith('/image') and rid not in used:doc.part.drop_rel(rid)
    clean=SUB/'Main_Text_CN_2009_2020_five_outcomes_clean.docx';doc.save(clean)
    highlighted=Document(clean)
    for p in highlighted.paragraphs:
        for r in p.runs:r.font.highlight_color=None
    for p in highlighted.paragraphs[32:39]:
        # Split runs at the first Chinese full stop, retaining all rPr/citations.
        remaining=p.text.index('。')+1
        for r in list(p.runs):
            if remaining<=0:break
            if len(r.text)<=remaining:r.font.highlight_color=WD_COLOR_INDEX.YELLOW;remaining-=len(r.text)
            else:
                tail=deepcopy(r._r);r._r.addnext(tail)
                from docx.text.run import Run
                tr=Run(tail,p);tr.text=r.text[remaining:];r.text=r.text[:remaining];r.font.highlight_color=WD_COLOR_INDEX.YELLOW;remaining=0
    high=SUB/'Main_Text_CN_2009_2020_five_outcomes_highlight.docx';highlighted.save(high)
    assert [p.text for p in Document(clean).paragraphs]==[p.text for p in Document(high).paragraphs]
    assert len(doc.element.xpath('.//m:oMath'))==28
    (AUD/'main_change_details.json').write_text(json.dumps(changes,ensure_ascii=False,indent=2),encoding='utf-8')
    (AUD/'main_ordinary_text_only.md').write_text('\n\n'.join(p.text for p in doc.paragraphs),encoding='utf-8')
    print('Main clean/highlight saved; 28 native math objects preserved.')
if __name__=='__main__':run()
