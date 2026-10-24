"""Formatting and native-OMML helpers, mechanically adapted from the verified latest SI builder."""
from environment import ROOT,SUB,SD,AUD
from pathlib import Path
import hashlib,json,re,sys,zipfile
from latex2mathml.converter import convert as latex_to_mathml
from mathml2omml import convert as mathml_to_omml
from lxml import etree
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from PIL import Image
MM_NS='http://schemas.openxmlformats.org/officeDocument/2006/math'
W_NS='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
OUT=SUB;QC=AUD
INPUTS={};TABLE_RECORD=[];MATH_RECORD=[];FIGURE_RECORD=[]
def track(path):
    path=Path(path);INPUTS[str(path)]=hashlib.sha256(path.read_bytes()).hexdigest();return path


def math_element(latex):
    try: mathml=latex_to_mathml(latex)
    except Exception:
        print('FAILED_LATEX',repr(latex),flush=True)
        raise
    xml=mathml_to_omml(mathml)
    # mathml2omml emits a wrong closing tag for the properties of overbars.
    # Repair only that observed malformed tag, never alter the equation tokens.
    xml=re.sub(r'(<m:groupChrPr>.*?)</m:groupChr>',r'\1</m:groupChrPr>',xml)
    if 'xmlns:m=' not in xml:
        xml=xml.replace('<m:oMath',f'<m:oMath xmlns:m="{MM_NS}" xmlns:w="{W_NS}"',1)
    try:
        node=etree.fromstring(xml.encode())
    except etree.XMLSyntaxError:
        print('FAILED_MATH',latex,xml,flush=True)
        raise
    # Hide absent sum/product limits so LibreOffice does not paint empty boxes.
    for nary in node.findall('.//{%s}nary'%MM_NS):
        pr=nary.find('{%s}naryPr'%MM_NS)
        if pr is None:pr=etree.SubElement(nary,'{%s}naryPr'%MM_NS)
        for tag in ['sub','sup']:
            slot=nary.find('{%s}%s'%(MM_NS,tag))
            if slot is None or not ''.join(slot.itertext()).strip():
                hide=etree.SubElement(pr,'{%s}%sHide'%(MM_NS,tag));hide.set('{%s}val'%MM_NS,'1')
    for group in list(node.findall('.//{%s}groupChr'%MM_NS)):
        chr_el=group.find('{%s}groupChrPr/{%s}chr'%(MM_NS,MM_NS))
        if chr_el is not None and chr_el.get('{%s}val'%MM_NS)=='¯':
            bar=etree.Element('{%s}bar'%MM_NS);pr=etree.SubElement(bar,'{%s}barPr'%MM_NS)
            pos=etree.SubElement(pr,'{%s}pos'%MM_NS);pos.set('{%s}val'%MM_NS,'top')
            expr=group.find('{%s}e'%MM_NS);bar.append(expr);group.getparent().replace(group,bar)
    for radical in node.findall('.//{%s}rad'%MM_NS):
        if radical.find('{%s}deg'%MM_NS) is None:
            pr=etree.Element('{%s}radPr'%MM_NS)
            hide=etree.SubElement(pr,'{%s}degHide'%MM_NS);hide.set('{%s}val'%MM_NS,'1')
            radical.insert(0,pr);radical.insert(1,etree.Element('{%s}deg'%MM_NS))
    for accent_limit in list(node.findall('.//{%s}limUpp'%MM_NS)):
        limit=accent_limit.find('{%s}lim'%MM_NS)
        if limit is not None and ''.join(limit.itertext())=='^':
            accent=etree.Element('{%s}acc'%MM_NS)
            pr=etree.SubElement(accent,'{%s}accPr'%MM_NS)
            char=etree.SubElement(pr,'{%s}chr'%MM_NS);char.set('{%s}val'%MM_NS,'̂')
            accent.append(accent_limit.find('{%s}e'%MM_NS))
            accent_limit.getparent().replace(accent_limit,accent)
    for run in node.findall('.//{%s}r'%MM_NS):
        # Explicit math fonts avoid inherited CJK/Latin fallback mixing in Office.
        word_pr=etree.Element('{%s}rPr'%W_NS)
        run.insert(1 if run.find('{%s}rPr'%MM_NS) is not None else 0,word_pr)
        fonts=etree.SubElement(word_pr,'{%s}rFonts'%W_NS)
        for key in ['ascii','hAnsi','eastAsia','cs']:fonts.set('{%s}%s'%(W_NS,key),'Cambria Math')
        token=run.find('{%s}t'%MM_NS)
        if token is not None and token.text=='∣':token.text='|'
        if token is not None and token.text=='~' and r'\sim' in latex:
            token.text='∼'
            style=run.find('{%s}rPr/{%s}sty'%(MM_NS,MM_NS))
            if style is not None:style.set('{%s}val'%MM_NS,'p')
    MATH_RECORD.append(dict(latex=latex,math_text=''.join(node.itertext()),status='CREATED'))
    return node


def text_run(p,text,size=None,bold=False):
    r=p.add_run(text)
    r.font.name='Times New Roman'; r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'SimSun')
    r.font.color.rgb=RGBColor(0,0,0);r.bold=bold
    if size: r.font.size=Pt(size)
    return r


def inline(p,text,size=None,bold=False):
    for part in re.split(r'(\$[^$]+\$)',text):
        if not part:continue
        if part.startswith('$') and part.endswith('$'):p._p.append(math_element(part[1:-1]))
        else:
            # Units and chemical subscripts are also native Office Math, not
            # Unicode pseudo-superscripts or raw TeX in prose/table headings.
            pattern=r'(μg(?:\s*m[−-]3|/m³)|PM(?:2\.5|₂\.₅)|°C)'
            for chunk in re.split(pattern,part):
                if not chunk:continue
                if chunk.startswith('μg'):
                    p._p.append(math_element(r'\mu\mathrm{g}\,\mathrm{m}^{-3}'))
                elif chunk in ['PM2.5','PM₂.₅']:
                    p._p.append(math_element(r'\mathrm{PM}_{2.5}'))
                elif chunk=='°C':p._p.append(math_element(r'\mathrm{°C}'))
                else:text_run(p,chunk.replace('**','').replace('`',''),size,bold)


def add_table(doc,lines):
    data=[[c.strip() for c in line.strip().strip('|').split('|')] for line in lines]
    data=[data[0]]+[r for r in data[1:] if not all(re.fullmatch(r'[:\- ]+',c or '-') for c in r)]
    n=len(data[0]);t=doc.add_table(rows=1,cols=n);t.autofit=False
    if n==4 and data[0][0]=='ID':widths=[51,20,69,40]
    elif n==7 and data[0][1:3]==['参考模型','联合模型']:widths=[17,26,43,18,18,26,32]
    elif n==7 and data[0][0]=='Rank': widths=[11,51,14,23,30,20,21]
    elif n==7:widths=[17,37,17,24,39,23,23]
    elif n==6 and 'SD水平或模型' in data[0]:widths=[18,42,18,27,47,28]
    elif n==6:widths=[23,22,34,34,33,34]
    elif n==8:widths=[24,18,23,23,23,23,23,23]
    elif n==5:widths=[34,22,44,47,33]
    elif n==4:widths=[37,25,66,52]
    elif n==3:widths=[35,67,78]
    elif n==2:widths=[40,140]
    else:widths=[180/n]*n
    for col,width in zip(t.columns,widths):col.width=Mm(width)
    for ri,row in enumerate(data):
        cells=t.rows[0].cells if ri==0 else t.add_row().cells
        for ci,(cell,value) in enumerate(zip(cells,row)):
            cell.width=Mm(widths[ci]);cell.vertical_alignment=WD_ALIGN_VERTICAL.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            margin=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),'65');el.set(qn('w:type'),'dxa');margin.append(el)
            tcpr.append(margin)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EEF3' if ri==0 else 'FFFFFF');tcpr.append(shade)
            p=cell.paragraphs[0];p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.05
            p.paragraph_format.keep_with_next=False;p.paragraph_format.keep_together=False
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER if ri==0 or ci==0 or len(value)<18 else WD_ALIGN_PARAGRAPH.LEFT
            inline(p,value,9 if n>5 else 9.5,ri==0)
        trpr=t.rows[ri]._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
    t._tbl.tblPr.append(borders)
    return t


def build():
    INPUTS.clear();TABLE_RECORD.clear();MATH_RECORD.clear();FIGURE_RECORD.clear()
    OUT.mkdir(parents=True,exist_ok=True)
    text=prepare_markdown()
    (OUT/'Supplementary_Information_CN.md').write_text(text,encoding='utf-8')
    equations=re.findall(r'```latex\n(.*?)\n```',text,re.S)
    (OUT/'supplement_equations.tex').write_text('% Frozen supplementary equations; native OMML in Word.\n\n'+'\n\n'.join(equations),encoding='utf-8')
    doc=Document();sec=doc.sections[0]
    sec.page_width=Mm(210);sec.page_height=Mm(297)
    sec.top_margin=Mm(16);sec.bottom_margin=Mm(16);sec.left_margin=Mm(15);sec.right_margin=Mm(15)
    for name in ['Normal','Title','Heading 1','Heading 2','Heading 3','Caption']:
        style=doc.styles[name];style.font.name='Times New Roman';style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'SimSun');style.font.color.rgb=RGBColor(0,0,0)
        style.font.size=Pt(10.5 if name=='Normal' else 18 if name=='Title' else 13 if name=='Heading 1' else 11)
        style.paragraph_format.space_before=Pt(8 if 'Heading' in name else 0);style.paragraph_format.space_after=Pt(5)
    for border in list(doc.styles.element.xpath('.//w:pBdr')):border.getparent().remove(border)
    doc.styles['Normal'].paragraph_format.line_spacing=1.12
    doc.styles['Normal'].paragraph_format.widow_control=True
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    lines=text.splitlines();i=0;figure_next=False
    while i<len(lines):
        line=lines[i].strip()
        if not line:i+=1;continue
        if line.startswith('<!-- FIGURE_PAGE'):
            doc.add_page_break();figure_next=True;i+=1;continue
        if line.startswith('<!--'):i+=1;continue
        if line.startswith('```latex'):
            block=[];i+=1
            while i<len(lines) and not lines[i].startswith('```'):block.append(lines[i]);i+=1
            latex='\n'.join(block)
            # A line ending in a comma or incomplete delimiters continues; otherwise
            # use separate native display paragraphs to avoid page-width overflow.
            if latex.lstrip().startswith(r'\widehat\alpha'):
                equation_lines=[''.join(block)]
            elif '\n' in latex:
                equation_lines=[];acc=''
                for value in block:
                    acc+=value.strip()
                    if acc.count('{')==acc.count('}') and acc.count('\\left')==acc.count('\\right') and not acc.endswith(('\\propto',',','\\')):
                        equation_lines.append(acc);acc=''
                if acc:equation_lines.append(acc)
            else:equation_lines=[latex]
            for eq_index,value in enumerate(equation_lines):
                p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_together=True
                p.paragraph_format.keep_with_next=eq_index<len(equation_lines)-1
                p._p.append(math_element(value))
            i+=1;continue
        if line.startswith('|'):
            table_lines=[]
            while i<len(lines) and lines[i].strip().startswith('|'):table_lines.append(lines[i].strip());i+=1
            add_table(doc,table_lines);continue
        img=re.match(r'!\[([^]]*)\]\(([^)]+)\)',line)
        if img:
            path=(OUT/Path(img.group(2))).resolve();track(path)
            with Image.open(path) as im:w,h=im.size
            width=180; height=width*h/w
            if height>200:width*=200/height;height=200
            p=doc.add_paragraph();p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(4)
            pic=p.add_run().add_picture(str(path),width=Mm(width));pic._inline.docPr.set('descr',img.group(1))
            FIGURE_RECORD.append(dict(figure=img.group(1),path=str(path),width_mm=width,height_mm=height,embedded=True))
            i+=1;continue
        if line.startswith('#'):
            level=len(line)-len(line.lstrip('#'));title=line.lstrip('# ').replace('、',' ').replace('，',' ').replace('：',' ').replace('|',' ')
            p=doc.add_paragraph(style='Title' if level==1 else 'Heading 1' if level==2 else 'Heading 2')
            inline(p,title);p.paragraph_format.keep_with_next=True
            i+=1;continue
        if line.startswith('Equation ('):
            p=doc.add_paragraph();p.paragraph_format.keep_with_next=True;inline(p,line.replace('Equation','补充公式'));i+=1;continue
        p=doc.add_paragraph();inline(p,line)
        if line.startswith('补充图S'):
            p.paragraph_format.keep_together=True;p.paragraph_format.space_after=Pt(6)
            for r in p.runs:r.font.size=Pt(10)
        i+=1
    doc.core_properties.title='AFIRE研究补充信息';doc.core_properties.author='';doc.core_properties.comments=''
    output=OUT/'Supplementary_Information_CN_2009_2020_five_outcomes.docx';doc.save(output)
    reopened=Document(output)
    health_tables=[t for t in reopened.tables if 'SD水平或模型' in ''.join(c.text for c in t.rows[0].cells)]
    assert [len(t.rows)-1 for t in health_tables]==[5,25,40,25,40], 'Saved DOCX lost full results or headers'
    with zipfile.ZipFile(output) as z:xml=etree.fromstring(z.read('word/document.xml'))
    ns={'w':W_NS,'m':MM_NS}
    bodytext=''.join(xml.xpath('//w:t/text()',namespaces=ns))
    assert not re.search(r'\\(?:beta|frac|sum|mu|alpha|rho)|\$|```',bodytext),'Raw TeX in Word text'
    assert len(xml.xpath('//m:oMath',namespaces=ns))==len(MATH_RECORD)
    assert len(doc.inline_shapes)==7
    audit=dict(status='STRUCTURAL_PASS_VISUAL_PENDING',tables_135_rows=[r for r in TABLE_RECORD if r.get('analysis_block')],main_result_rows=135,n_tables=len(doc.tables),native_math_objects=len(MATH_RECORD),embedded_figures=FIGURE_RECORD,raw_latex_in_word=False,source_hashes_unchanged=all(hashlib.sha256(Path(p).read_bytes()).hexdigest()==v for p,v in INPUTS.items()),table_source_checks=TABLE_RECORD)
    (QC/'supplement_build_check.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
    (QC/'supplement_math_objects.json').write_text(json.dumps(MATH_RECORD,ensure_ascii=False,indent=2),encoding='utf-8')
    (QC/'supplement_input_hashes.json').write_text(json.dumps(INPUTS,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({k:v for k,v in audit.items() if k not in ['embedded_figures','table_source_checks']},ensure_ascii=False,indent=2))
    return output



